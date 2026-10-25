"""PDFs found anywhere: pictures, ruled tables, lists and titles that pdftotext alone breaks."""
import io

from docx import Document
from docx.shared import Cm
from PIL import Image, ImageDraw

from app.doc.extract import from_pdf, from_text, pdf_cleanup
from app.doc.office import finalize


def test_numbers_and_bullets_printed_apart_go_back_to_their_items():
    text = "Les étapes sont :\n1)\n2)\n3)\n\nOn envoie le bloc.\nChaque nœud le vérifie et\nl'accepte.\nOn passe au suivant.\n-\nSuivre la trésorerie.\n"
    out = [ln for ln in pdf_cleanup(text).split("\n") if ln.strip()]
    assert out[1:4] == ["1) On envoie le bloc.", "2) Chaque nœud le vérifie et l'accepte.", "3) On passe au suivant."]
    assert "- Suivre la trésorerie." in out


def test_section_number_at_the_top_of_pages_is_not_a_running_header():
    titles = ["Introduction", "Transactions", "Timestamp Server", "Proof-of-Work", "Network"]
    pages = [f"{n}.\n{t}\nCe que dit la partie sur {t.lower()}." for n, t in enumerate(titles, 1)]
    out = pdf_cleanup("\f".join(pages))
    assert "2. Transactions" in out and "5. Network" in out


def test_capital_title_on_three_lines_at_the_top_of_a_page_is_one_title():
    text = pdf_cleanup("Fin de l'introduction.\fPARTIE I : PRESENTATION\nGLOBALE DU CADRE\nINSTITUTIONNEL\nCette partie intitulée présentation du cadre comporte deux\nchapitres. Le premier présente la CPC.")
    raws = from_text(text, source="pdf")
    assert raws[1].text == "PARTIE I : PRESENTATION GLOBALE DU CADRE INSTITUTIONNEL" and raws[1].new_page
    assert raws[2].text.startswith("Cette partie") and raws[2].text.endswith("présente la CPC.")


def test_sentence_going_on_over_the_page_stays_one_paragraph():
    raws = from_text(pdf_cleanup("Le directeur des opérations établit et négocie l'affectation du\fpersonnel aux différentes fonctions."), source="pdf")
    assert len(raws) == 1 and not raws[0].new_page


def test_pictures_and_ruled_tables_of_a_pdf_come_back(tmp_path):
    img = Image.new("RGB", (600, 360), "white")
    d = ImageDraw.Draw(img)
    d.rectangle((20, 20, 580, 340), outline="navy", width=6)
    d.ellipse((120, 60, 480, 300), fill="orange")
    buf = io.BytesIO()
    img.save(buf, "PNG")
    doc = Document()
    doc.add_paragraph("Le schéma ci-dessous présente le circuit des documents dans le service.")
    doc.add_picture(io.BytesIO(buf.getvalue()), width=Cm(12))
    doc.add_paragraph("Figure 1 : circuit des documents")
    doc.add_paragraph("Le tableau suivant résume les effectifs du service par année.")
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    for r, row in enumerate([["Année", "Agents", "Dossiers"], ["2021", "12", "340"], ["2022", "15", "410"]]):
        for c, v in enumerate(row):
            table.cell(r, c).text = v
    doc.add_paragraph("Ces chiffres montrent une hausse régulière de l'activité.")
    src = tmp_path / "in.docx"
    doc.save(src)
    pdf = tmp_path / "in.pdf"
    finalize(src, None, pdf)

    raws = from_pdf(pdf, tmp_path / "images")
    figures = [r for r in raws if r.image]
    tables = [r for r in raws if r.rows]
    assert len(figures) == 1 and (tmp_path / "images" / figures[0].image).is_file()
    assert tables and tables[0].rows[0] == ["Année", "Agents", "Dossiers"] and tables[0].rows[2] == ["2022", "15", "410"]
    texts = [r.text for r in raws if r.text]
    assert "Figure 1 : circuit des documents" in texts
    assert not any("2021" in t for t in texts)  # the cells are in the table, not left in the text
    order = [("fig" if r.image else "tab" if r.rows else r.text[:10]) for r in raws]
    assert order.index("fig") < order.index("Figure 1 :") < order.index("tab")
