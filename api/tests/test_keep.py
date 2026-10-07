"""A Word file that already has its cover page / sommaire / numbers is kept; only what is missing is added."""
import base64
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app import config, db, fapshi
from app.doc import keep


def cover_docx(tmp_path, numbered=False, toc=False):
    doc = Document()
    for line in ("RÉPUBLIQUE DU CAMEROUN", "UNIVERSITÉ D'EBOLOWA", "RAPPORT DE STAGE", "Par : Jean Dupont", "Matricule : 22X0070",
                 "Année académique : 2026-2027"):
        doc.add_paragraph(line)
    doc.add_page_break()
    if toc:
        doc.add_paragraph("Sommaire")
    doc.add_heading("Introduction", 1)
    doc.add_paragraph("Texte de l'introduction. " * 20)
    doc.add_heading("Conclusion", 1)
    doc.add_paragraph("Texte de la conclusion.")
    if numbered:
        keep_field = doc.sections[0].footer.paragraphs[0]
        from app.doc.render import _field
        _field(keep_field, "PAGE")
    path = tmp_path / "in.docx"
    doc.save(path)
    return path


def test_inspect_sees_the_cover_and_what_is_missing(tmp_path):
    info = keep.inspect(cover_docx(tmp_path))
    assert info["cover"] is True and info["page_numbers"] is False and info["toc"] is False and info["headings"] == 2


def test_inspect_sees_existing_numbers(tmp_path):
    assert keep.inspect(cover_docx(tmp_path, numbered=True))["page_numbers"] is True


def test_touch_adds_numbers_and_leaves_the_cover_alone(tmp_path):
    src = cover_docx(tmp_path)
    out = tmp_path / "out.docx"
    assert keep.touch(src, out, page_numbers=True) == ["numéros de page"]
    after = Document(str(out))
    assert [p.text for p in after.paragraphs[:6]] == [p.text for p in Document(str(src)).paragraphs[:6]]
    assert len(after.sections) == 2  # the cover is a section of its own, counted but without a number
    assert not keep._section_has_numbers(after.sections[0]) and keep._section_has_numbers(after.sections[1])
    assert keep.inspect(out)["page_numbers"] is True
    assert keep.touch(out, tmp_path / "again.docx", page_numbers=True) == []  # nothing to add the second time


def test_existing_footer_text_is_never_overwritten(tmp_path):
    doc = Document()
    doc.add_paragraph("Texte")
    doc.sections[0].footer.paragraphs[0].text = "Confidentiel"
    src = tmp_path / "f.docx"
    doc.save(src)
    keep.touch(src, tmp_path / "o.docx", page_numbers=True)
    assert Document(str(tmp_path / "o.docx")).sections[0].footer.paragraphs[0].text == "Confidentiel"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    for mod in (config, db):
        monkeypatch.setattr(mod, "DATA_DIR", tmp_path)
        monkeypatch.setattr(mod, "DB_PATH", tmp_path / "t.sqlite3")
    monkeypatch.setattr(config, "FAPSHI_MODE", "mock")
    monkeypatch.setattr(fapshi, "MOCK_DELAY_S", 0)
    from app.main import app
    with TestClient(app) as c:
        yield c


def test_import_keeps_a_file_that_already_has_a_cover_and_the_render_keeps_it(client, tmp_path):
    data = base64.b64encode(cover_docx(tmp_path).read_bytes()).decode()
    doc = client.post("/documents", json={"filename": "r.docx", "data": data}).json()
    assert doc["mode"] == "keep" and doc["plan"] == {"cover": "keep", "toc": "none", "numbers": "add"}
    assert doc["confirmed"] is False and doc["analysis"]["existing"]["cover"] is True
    assert doc["meta"]["existing"]["cover"] is True
    rendered = client.post(f"/documents/{doc['id']}/render").json()
    assert rendered["render"]["pages"] >= 2
    word = client.get(f"/documents/{doc['id']}")  # still readable
    assert word.status_code == 200
    # the user can switch to the full Paginya layout, and back
    assert client.put(f"/documents/{doc['id']}", json={"mode": "rebuild"}).json()["mode"] == "rebuild"
    back = client.put(f"/documents/{doc['id']}", json={"mode": "keep", "plan": {"numbers": "redo", "toc": "add"}, "confirm": True}).json()
    assert back["plan"] == {"cover": "keep", "toc": "add", "numbers": "redo"} and back["confirmed"] is True


def test_plain_file_without_cover_still_gets_the_normal_layout(client, tmp_path):
    doc = Document()
    doc.add_heading("Titre", 1)
    doc.add_paragraph("Un texte simple.")
    path = tmp_path / "p.docx"
    doc.save(path)
    data = base64.b64encode(path.read_bytes()).decode()
    made = client.post("/documents", json={"filename": "p.docx", "data": data}).json()
    assert made["mode"] == "rebuild"


# --- the plan, element by element -------------------------------------------------------------------

from app.doc.analysis import analyze, clean_plan, default_plan  # noqa: E402
from app.doc.exam_scan import exam_signals, split_type  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Pt  # noqa: E402


def report_docx(tmp_path, name="r.docx", plain_titles=True, prelim=True, numbers=False):
    d = Document()
    for line in ("RÉPUBLIQUE DU CAMEROUN", "UNIVERSITÉ D'EBOLOWA", "RAPPORT DE STAGE", "Par : Jean Dupont", "Matricule : 22X0070",
                 "Année académique : 2026-2027"):
        d.add_paragraph(line)
    d.add_page_break()

    def title(text, level=1):
        if plain_titles:
            run = d.add_paragraph().add_run(text)
            run.bold = True
            run.font.size = Pt(14 if level == 1 else 12)
        else:
            d.add_heading(text, level)

    if prelim:
        title("REMERCIEMENTS")
        d.add_paragraph("Merci. " * 30)
        d.add_page_break()
    title("INTRODUCTION")
    d.add_paragraph("Texte. " * 60)
    d.add_page_break()
    title("CHAPITRE I : PRÉSENTATION")
    title("1. Historique", 2)
    d.add_paragraph("Texte. " * 60)
    title("CONCLUSION")
    d.add_paragraph("Fin.")
    if numbers:
        from app.doc.render import _field
        _field(d.sections[0].footer.paragraphs[0], "PAGE")
    path = tmp_path / name
    d.save(path)
    return path


def formats(path):
    return [(s._sectPr.find(qn("w:pgNumType")).get(qn("w:fmt")) if s._sectPr.find(qn("w:pgNumType")) is not None else None,
             keep._section_has_numbers(s)) for s in Document(str(path)).sections]


def test_titles_are_found_without_word_heading_styles(tmp_path):
    info = keep.inspect(report_docx(tmp_path))
    assert info["headings"] == 0 and info["titles"] == 5 and info["intro"] is True


def test_numbers_follow_the_cameroonian_scheme(tmp_path):
    out = tmp_path / "o.docx"
    assert "numéros de page" in keep.touch(report_docx(tmp_path), out, numbers="add")
    # cover (roman, no number) / preliminary pages (roman) / body (arabic from the Introduction)
    assert formats(out) == [("lowerRoman", False), ("lowerRoman", True), ("decimal", True)]
    assert Document(str(out)).sections[2]._sectPr.find(qn("w:pgNumType")).get(qn("w:start")) == "1"


def test_redo_replaces_numbers_everywhere_on_the_old_scheme(tmp_path):
    src = report_docx(tmp_path, numbers=True)
    assert keep.inspect(src)["numbering"] == "arabic" and keep.inspect(src)["numbers_on_cover"] is True
    out = tmp_path / "o.docx"
    keep.touch(src, out, numbers="redo")
    assert keep.inspect(out)["numbering"] == "roman_arabic" and keep.inspect(out)["numbers_on_cover"] is False
    assert keep.touch(src, tmp_path / "k.docx", numbers="keep") == []  # "keep" never touches them
    assert formats(tmp_path / "k.docx") == formats(src)


def test_toc_is_added_to_titles_that_have_no_heading_style_and_redone_when_asked(tmp_path):
    out = tmp_path / "o.docx"
    assert "sommaire" in keep.touch(report_docx(tmp_path), out, toc="add")
    texts = [p.text for p in Document(str(out)).paragraphs]
    assert "Sommaire" in texts and "[[TOC:2]]" in texts
    again = tmp_path / "again.docx"
    keep.touch(out, again, toc="redo")  # the marker is not a Word sommaire yet: nothing is duplicated
    assert [p.text for p in Document(str(again)).paragraphs].count("[[TOC:2]]") == 1


def test_cover_can_be_redone_with_paginya_cover(tmp_path):
    from app import render as cover_render
    svg = '<svg xmlns="http://www.w3.org/2000/svg" width="595" height="842"><rect width="595" height="842" fill="#eee"/></svg>'
    png = cover_render.svg_to_png(svg, 60)
    out = tmp_path / "o.docx"
    done = keep.touch(report_docx(tmp_path), out, cover="redo", cover_png=png, numbers="add")
    assert "page de garde" in done
    after = Document(str(out))
    assert "RAPPORT DE STAGE" not in [p.text for p in after.paragraphs]  # the old cover is gone
    assert len(after.inline_shapes) == 1
    assert formats(out)[0] == ("lowerRoman", False)  # the new cover is counted, not numbered


def test_exam_is_recognised_from_its_text_and_the_ai_reply():
    assert exam_signals("Épreuve de maths\nExercice 1 (5 pts)\n1) Calcule (2 pts)\nDurée : 2h")["likely"] is True
    assert exam_signals("INTRODUCTION\nCe rapport présente le stage.")["likely"] is False
    assert split_type("TYPE: epreuve\nENTETE\necole: Lycée") == ("epreuve", "ENTETE\necole: Lycée")
    assert split_type("TYPE: document\nBonjour")[0] == "document"
    assert split_type("Bonjour") == ("", "Bonjour")


def test_analysis_proposes_a_plan_and_a_hint_about_numbering():
    existing = {"cover": True, "toc": False, "page_numbers": True, "numbering": "arabic", "numbers_on_cover": False, "intro": True, "titles": 8}
    report = analyze("rapport_stage", {"words": 5000, "headings": 8}, [{"type": "paragraph", "text": "Texte"}], {"cover": True, "toc": True, "page_numbers": True}, existing, True)
    assert report["plan"] == {"cover": "keep", "toc": "add", "numbers": "keep"} and "romains" in report["hint"]
    assert report["kind_label"] == "un rapport de stage" and report["exam"] is False
    assert clean_plan({"toc": "redo", "cover": "bogus"}, report["plan"]) == {"cover": "keep", "toc": "redo", "numbers": "keep"}
    exam = analyze("document", {}, [{"type": "paragraph", "text": "Épreuve de maths\nExercice 1 (5 pts)\n1) Calcule (2 pts)\nDurée : 2h"}], {}, None, False)
    assert exam["kind"] == "epreuve" and exam["exam"] is True


def test_nothing_is_rendered_before_the_user_confirms_and_new_documents_are_unconfirmed(client):
    made = client.post("/documents", json={"text": "INTRODUCTION\nBonjour le monde, ceci est un texte assez long pour un test."}).json()
    assert made["confirmed"] is False and made["render"] is None and made["analysis"]["docx"] is False


def test_scan_without_kind_decides_by_itself(client, monkeypatch):
    import app.doc_routes as routes
    monkeypatch.setattr(routes, "straighten", lambda raw: (raw, False))
    monkeypatch.setattr(routes, "ocr_local", lambda jpeg: ("Épreuve de maths\nExercice 1 (5 pts)\n1) Calcule (2 pts)\nDurée : 2h", 90.0))
    data = base64.b64encode(b"x").decode()
    got = client.post("/scans", json={"data": data}).json()
    assert got["is_exam"] is True and "exam" in got and "Exercice 1" in got["text"]
    monkeypatch.setattr(routes, "ocr_local", lambda jpeg: ("INTRODUCTION\nBonjour tout le monde.", 90.0))
    got = client.post("/scans", json={"data": data}).json()
    assert got["is_exam"] is False and "exam" not in got


def test_full_flow_confirm_then_render_with_the_cameroonian_numbers(client, tmp_path):
    import subprocess

    data = base64.b64encode(report_docx(tmp_path, numbers=True).read_bytes()).decode()
    doc = client.post("/documents", json={"filename": "r.docx", "data": data}).json()
    assert doc["mode"] == "keep" and doc["plan"]["numbers"] == "keep" and doc["analysis"]["hint"]
    done = client.put(f"/documents/{doc['id']}", json={"plan": {"numbers": "redo", "toc": "add"}, "confirm": True}).json()
    assert done["confirmed"] is True and done["plan"] == {"cover": "keep", "toc": "add", "numbers": "redo"}
    rendered = client.post(f"/documents/{doc['id']}/render").json()
    assert rendered["render"]["pages"] >= 5
    from app.doc_routes import doc_dir
    pdf = doc_dir(doc["id"]) / "build" / "final.pdf"
    feet = []
    for page in range(1, rendered["render"]["pages"] + 1):
        text = subprocess.run(["pdftotext", "-f", str(page), "-l", str(page), "-layout", str(pdf), "-"], capture_output=True, text=True).stdout
        lines = [ln.strip() for ln in text.splitlines() if ln.strip()]
        feet.append(lines[-1] if lines else "")
    assert feet[0] != "1" and feet[1] == "ii"  # cover: no number, then the sommaire is "ii"
    assert "1" in feet  # the body restarts at 1 on the Introduction
