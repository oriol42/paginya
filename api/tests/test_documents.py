import base64
import io
import shutil
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Pt
from fastapi.testclient import TestClient
from PIL import Image

from app import config, db, fapshi
from app.doc.detect import detect
from app.doc.extract import from_docx, from_text

FIXTURE = Path(__file__).parent / "fixtures" / "rapport_brut.txt"


@pytest.fixture()
def client(tmp_path, monkeypatch):
    for mod in (config, db):
        monkeypatch.setattr(mod, "DATA_DIR", tmp_path)
        monkeypatch.setattr(mod, "DB_PATH", tmp_path / "t.sqlite3")
    monkeypatch.setattr(config, "FAPSHI_MODE", "mock")
    monkeypatch.setattr(fapshi, "MOCK_DELAY_S", 0)
    from app.main import app
    with TestClient(app) as c:
        yield c


def headings(blocks):
    return [(b["text"], b["level"]) for b in blocks if b["type"] == "heading"]


def test_plain_text_structure():
    res = detect(from_text(FIXTURE.read_text()))
    h = dict(headings(res["blocks"]))
    assert h["PREMIÈRE PARTIE : PRÉSENTATION DE LA STRUCTURE D'ACCUEIL"] == 1
    assert h["CHAPITRE I : HISTORIQUE ET MISSIONS"] == 2
    assert h["I. Historique"] == 3
    assert h["1.1 Tri et classement des archives"] == 3  # same depth as "I." in another chapter
    types = [b["type"] for b in res["blocks"]]
    assert "SOMMAIRE" not in [b.get("text") for b in res["blocks"]]  # old TOC is dropped
    # Table caption moved above its table, table rebuilt from tab-separated lines
    i = types.index("table")
    assert res["blocks"][i - 1]["type"] == "caption" and res["blocks"][i - 1]["of"] == "table"
    assert res["blocks"][i]["rows"][0] == ["Poste", "Effectif", "Rôle"]
    lists = [b for b in res["blocks"] if b["type"] == "list"]
    assert sum(b["ordered"] for b in lists) == 3 and sum(not b["ordered"] for b in lists) == 6
    assert any(b.get("role") == "sigle" and b["term"] == "COBAC" for b in res["blocks"])
    assert any(b.get("role") == "dedicace" for b in res["blocks"])


def test_numbered_line_alone_is_a_heading_but_a_run_is_a_list():
    text = "1. Contexte de l'étude\nLe Cameroun compte dix régions et une grande diversité culturelle qui influence fortement les pratiques documentaires des institutions publiques et privées.\n\n1. Premier point\n2. Deuxième point\n"
    blocks = detect(from_text(text))["blocks"]
    assert blocks[0]["type"] == "heading"
    assert [b["type"] for b in blocks[2:]] == ["list", "list"]


def make_docx(path: Path) -> None:
    d = Document()
    d.add_heading("Introduction", level=1)
    d.add_paragraph("Texte de l'introduction assez long pour être un paragraphe normal du document.")
    d.add_heading("Méthodologie", level=1)
    p = d.add_paragraph()
    r = p.add_run("Collecte des données")
    r.bold = True
    d.add_paragraph("Premier élément", style="List Number")
    d.add_paragraph("Second élément", style="List Number")
    d.add_paragraph("Une puce", style="List Bullet")
    t = d.add_table(rows=2, cols=2)
    t.cell(0, 0).text, t.cell(0, 1).text = "Région", "Effectif"
    t.cell(1, 0).text, t.cell(1, 1).text = "Centre", "120"
    buf = io.BytesIO()
    Image.new("RGB", (400, 200), "#0E9F6E").save(buf, format="PNG")
    buf.seek(0)
    d.add_picture(buf)
    d.add_paragraph("Figure 1 : Carte de la région")
    for par in d.paragraphs:
        for run in par.runs:
            run.font.size = run.font.size or Pt(12)
    d.save(path)


def test_docx_extraction(tmp_path):
    src = tmp_path / "in.docx"
    make_docx(src)
    raws = from_docx(src, tmp_path / "img")
    blocks = detect(raws)["blocks"]
    kinds = [(b["type"], b.get("text", "")[:20]) for b in blocks]
    assert ("heading", "Introduction") in kinds
    assert ("heading", "Collecte des données") in kinds  # bold short line
    lists = [b for b in blocks if b["type"] == "list"]
    assert [b["ordered"] for b in lists] == [True, True, False]
    assert any(b["type"] == "table" for b in blocks)
    fig = next(b for b in blocks if b["type"] == "figure")
    assert (tmp_path / "img" / fig["image"]).is_file()
    # figure caption stays below the figure
    i = blocks.index(fig)
    assert blocks[i + 1]["type"] == "caption" and blocks[i + 1]["of"] == "figure"


def test_full_flow_text_to_paid_download(client):
    doc = client.post("/documents", json={"text": FIXTURE.read_text()}).json()
    assert doc["meta"]["kind"] == "rapport_stage"
    r = client.post(f"/documents/{doc['id']}/render")
    assert r.status_code == 200, r.text
    info = r.json()
    assert info["render"]["pages"] >= 10
    assert info["product"] == "document_court" and info["amount"] == 550
    assert client.get(f"/documents/{doc['id']}/pages/1.webp").content[8:12] == b"WEBP"
    assert client.get(f"/orders/{doc['id']}/file.pdf").status_code == 402

    client.post(f"/orders/{doc['id']}/pay", json={"phone": "670000000"})
    assert client.get(f"/orders/{doc['id']}").json()["status"] == "PAID"
    pdf = client.get(f"/orders/{doc['id']}/file.pdf")
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    docx = client.get(f"/orders/{doc['id']}/file.docx")
    xml = zipfile.ZipFile(io.BytesIO(docx.content)).read("word/document.xml").decode()
    assert "Répartition du personnel" in xml and "[[TOC" not in xml  # markers replaced by real indexes


def test_edit_blocks_and_style(client):
    doc = client.post("/documents", json={"text": FIXTURE.read_text()}).json()
    blocks = doc["blocks"]
    blocks[0]["hidden"] = True
    r = client.put(f"/documents/{doc['id']}", json={"blocks": blocks, "style": {"theme": "moderne", "color": "#1E3A8A"}})
    assert r.status_code == 200
    assert r.json()["style"]["theme"] == "moderne" and r.json()["blocks"][0]["hidden"] is True
    bad = client.put(f"/documents/{doc['id']}", json={"style": {"theme": "hack"}})
    assert bad.json()["style"]["theme"] == "moderne"


def test_upload_docx(client, tmp_path):
    src = tmp_path / "in.docx"
    make_docx(src)
    r = client.post("/documents", json={"filename": "rapport.docx", "data": base64.b64encode(src.read_bytes()).decode()})
    assert r.status_code == 200, r.text
    assert any(b["type"] == "figure" for b in r.json()["blocks"])
    assert client.post(f"/documents/{r.json()['id']}/render").status_code == 200


def test_old_word_formats_are_converted(client, tmp_path):
    """.doc / .odt / .rtf: LibreOffice saves them as .docx on arrival, then they are Word files like any other."""
    rtf = (r"{\rtf1\ansi\deff0 {\fonttbl{\f0 Times New Roman;}}\f0\fs24 "
           r"{\b INTRODUCTION}\par Ce rapport pr\'e9sente le stage effectu\'e9 \'e0 Douala pendant trois mois.\par "
           r"{\b I- Pr\'e9sentation de l'entreprise}\par L'entreprise a \'e9t\'e9 cr\'e9\'e9e en 2004 et compte cent employ\'e9s.\par }").encode()
    r = client.post("/documents", json={"filename": "rapport.rtf", "data": base64.b64encode(rtf).decode()})
    assert r.status_code == 200, r.text
    doc = r.json()
    assert [b["text"] for b in doc["blocks"] if b["type"] == "heading"][:2] == ["INTRODUCTION", "I- Présentation de l'entreprise"]
    assert doc["analysis"]["docx"] is True


def photo_of_page(tmp_path) -> bytes:
    """A white page with text, photographed at an angle on a dark table."""
    import cv2
    import numpy as np
    from PIL import ImageDraw
    from PIL import ImageFont
    page = Image.new("RGB", (600, 850), "white")
    d = ImageDraw.Draw(page)
    font = ImageFont.truetype(str(Path(__file__).parent.parent / "fonts" / "tinos-400.ttf"), 26)
    d.text((50, 50), "CHAPITRE I : PRÉSENTATION", fill="black", font=font)
    for i in range(10):
        d.text((50, 120 + i * 60), f"Ligne de texte numéro {i + 1} du rapport", fill="black", font=font)
    src = np.array(page)[:, :, ::-1]
    table = np.full((1200, 1000, 3), (40, 50, 60), np.uint8)
    corners = np.float32([[0, 0], [599, 0], [599, 849], [0, 849]])
    target = np.float32([[180, 150], [800, 210], [760, 1080], [130, 1010]])
    warped = cv2.warpPerspective(src, cv2.getPerspectiveTransform(corners, target), (1000, 1200),
                                 dst=table, borderMode=cv2.BORDER_TRANSPARENT)
    ok, jpg = cv2.imencode(".jpg", warped)
    return jpg.tobytes()


def test_straighten_finds_the_sheet(tmp_path):
    from app.doc.scan import straighten
    jpeg, found = straighten(photo_of_page(tmp_path))
    assert found
    img = Image.open(io.BytesIO(jpeg))
    assert 0.6 < img.width / img.height < 0.8  # upright A4-ish page, table cropped away


def test_handwriting_requires_consent(client, monkeypatch, tmp_path):
    monkeypatch.setenv("OCR_PROVIDER", "mock")
    data = base64.b64encode(photo_of_page(tmp_path)).decode()
    assert client.post("/scans", json={"data": data, "handwriting": True}).status_code == 400
    r = client.post("/scans", json={"data": data, "handwriting": True, "consent": True})
    assert r.status_code == 200 and r.json()["straightened"] is True
    assert client.get(f"/scans/{r.json()['id']}.jpg").status_code == 200


@pytest.mark.skipif(shutil.which("tesseract") is None, reason="tesseract not installed")
def test_printed_page_read_locally_with_tesseract(client, tmp_path):
    data = base64.b64encode(photo_of_page(tmp_path)).decode()
    r = client.post("/scans", json={"data": data})  # no consent needed: stays on our server
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["engine"] == "tesseract"
    assert "CHAPITRE I" in body["text"] and "numéro 7" in body["text"] and body["confidence"] > 80


def test_ocr_rejoins_wrapped_lines():
    from app.doc.scan import _tsv_to_text
    rows = ["level\tpage_num\tblock_num\tpar_num\tline_num\tword_num\tleft\ttop\twidth\theight\tconf\ttext"]
    def w(line, left, width, text, par=1):
        rows.append(f"5\t1\t1\t{par}\t{line}\t1\t{left}\t{line * 30}\t{width}\t20\t95\t{text}")
    w(1, 100, 300, "INTRODUCTION")  # short title line
    w(2, 100, 890, "Comme le disait le philosophe, la théorie sans")  # reaches the right margin
    w(3, 100, 400, "pratique est inutile.")
    text, conf = _tsv_to_text("\n".join(rows))
    assert text == "INTRODUCTION\n\nComme le disait le philosophe, la théorie sans pratique est inutile."


def test_before_view_and_delete(client):
    doc = client.post("/documents", json={"text": FIXTURE.read_text()}).json()
    assert client.post(f"/documents/{doc['id']}/before").json()["pages"] >= 1
    assert client.get(f"/documents/{doc['id']}/before/1.png").status_code == 200
    assert doc["meta"]["changes"]
    assert client.delete(f"/documents/{doc['id']}").json()["deleted"] is True
    assert client.get(f"/documents/{doc['id']}").status_code == 404


def test_new_themes_render(client):
    doc = client.post("/documents", json={"text": FIXTURE.read_text()}).json()
    for theme in ("universitaire", "elegant", "corporate"):
        client.put(f"/documents/{doc['id']}", json={"style": {"theme": theme}})
        assert client.post(f"/documents/{doc['id']}/render").status_code == 200


def test_cleanup_after_7_days_and_purge_after_a_year(client):
    import time as _t
    from app.doc_routes import cleanup, doc_dir
    doc = client.post("/documents", json={"text": FIXTURE.read_text()}).json()
    assert doc_dir(doc["id"]).exists()
    cleanup(now=_t.time() + 8 * 86400)
    assert not doc_dir(doc["id"]).exists()
    assert client.get(f"/documents/{doc['id']}").status_code == 404
    cleanup(now=_t.time() + 400 * 86400)
    assert db.get_order(doc["id"]) is None


def test_survives_a_server_restart(client, tmp_path):
    """Hugging Face wipes the disk at every restart: the original, its images and the renders must come back."""
    src = tmp_path / "in.docx"
    make_docx(src)
    doc = client.post("/documents", json={"filename": "rapport.docx", "data": base64.b64encode(src.read_bytes()).decode()}).json()
    client.post(f"/documents/{doc['id']}/render")
    client.post(f"/orders/{doc['id']}/pay", json={"phone": "670000000"})
    assert client.get(f"/orders/{doc['id']}").json()["status"] == "PAID"

    shutil.rmtree(tmp_path / "docs")  # the restart: only the database and durable storage are left
    assert client.get(f"/documents/{doc['id']}/pages/1.webp").content[8:12] == b"WEBP"
    fig = next(b for b in doc["blocks"] if b["type"] == "figure")
    assert client.get(f"/documents/{doc['id']}/images/{fig['image']}").status_code == 200

    shutil.rmtree(tmp_path / "docs")
    assert client.get(f"/orders/{doc['id']}/file.pdf").content.startswith(b"%PDF")
    assert client.post(f"/documents/{doc['id']}/before").json()["pages"] >= 1

    client.delete(f"/documents/{doc['id']}")
    assert not [p for p in (tmp_path / "objects").rglob("*") if p.is_file()]  # right to erasure: durable copies deleted too


def test_form_survives_a_server_restart(client, tmp_path):
    from tests.test_forms import LETTER  # noqa: PLC0415

    f = client.post("/forms", json={"kind": "lettre", "data": LETTER}).json()
    shutil.rmtree(tmp_path / "forms")
    assert client.get(f"/forms/{f['id']}/pages/1.webp").content[8:12] == b"WEBP"


def test_a_conversion_changes_the_format_and_nothing_else(client, tmp_path):
    import base64
    from docx import Document

    d = Document()
    d.add_paragraph("Mon titre")
    d.add_paragraph("Un paragraphe de texte. " * 10)
    path = tmp_path / "c.docx"
    d.save(path)
    data = base64.b64encode(path.read_bytes()).decode()
    view = client.post("/documents", json={"filename": "c.docx", "data": data, "convert": True}).json()
    assert view["convert"] == "pdf" and view["confirmed"] is True and view["mode"] == "keep"
    assert view["plan"] == {"cover": "keep", "toc": "keep", "numbers": "keep"}
    assert client.post("/documents", json={"filename": "c.txt", "data": data, "convert": True}).status_code == 400


def test_numbered_titles_keep_one_level_per_numbering():
    """"I) / 1) / 2)": the items under a roman title all sit one level below it (the model must not move them)."""
    long = "Cette structure occupe une place importante dans le tissu économique de la région du Centre. " * 3
    text = "\n\n".join(["Introduction", long, "I) Les causes lointaines", "1) Les rivalités coloniales", long, "2) La course aux armements", long,
                         "II) Les causes immédiates", "1) L'attentat de Sarajevo", long, "2) L'engrenage des alliances", long, "Conclusion", long])
    h = dict(headings(detect(from_text(text))["blocks"]))
    assert h["I) Les causes lointaines"] == h["II) Les causes immédiates"] == 2
    assert {h[k] for k in h if k[0].isdigit()} == {3}


def test_body_is_never_hidden_as_a_cover_page(tmp_path):
    """A cover followed by titles the engine does not know by name ("A/ Objectif du TP") used to be read as one
    long cover page: every short line of the document was hidden."""
    doc = Document()
    for line in ("RÉPUBLIQUE DU CAMEROUN", "INSTITUT UNIVERSITAIRE DE TECHNOLOGIE", "COMPTE RENDU DE TP", "Rédigé et présenté par :",
                 "MBIDA Kevin", "Sous la supervision de :", "M. TAKAM", "Année académique 2025-2026"):
        doc.add_paragraph(line)
    for title in ("A/ Objectif du TP", "B/ Matériel utilisé", "C/ Manipulations"):
        doc.add_paragraph(title).runs[0].bold = True
        doc.add_paragraph("Nous avons configuré le réseau local de la salle avec trois commutateurs et un routeur.")
    path = tmp_path / "tp.docx"
    doc.save(path)
    blocks = detect(from_docx(path, tmp_path / "img"))["blocks"]
    shown = [b["text"] for b in blocks if not b.get("hidden") and b.get("text")]
    assert [t for t in shown if t[1:2] == "/"] == ["A/ Objectif du TP", "B/ Matériel utilisé", "C/ Manipulations"]
    assert sum("Nous avons configuré" in t for t in shown) == 3
    assert all(b.get("hidden") for b in blocks if b.get("text") in ("MBIDA Kevin", "COMPTE RENDU DE TP"))


def test_announced_chapters_are_a_list_and_typed_lines_are_cleaned():
    long = "Le service dispose de cinq ordinateurs, de deux imprimantes et d'une connexion internet instable. " * 3
    text = "\n\n".join(["INTRODUCTION", "Ce rapport s'articule autour de trois chapitres :", "Chapitre 1 : Présentation", "Chapitre 2 : Déroulement",
                         "Chapitre 3 : Bilan", "__________________________", "CHAPITRE 1 : PRESENTATION", long, "Plan du chapitre……………………",
                         "CHAPITRE 2 : DEROULEMENT", long, "CHAPITRE 3 : BILAN", long])
    blocks = detect(from_text(text))["blocks"]
    chapters = [b["text"] for b in blocks if b["type"] == "heading" and b["text"].upper().startswith("CHAPITRE")]
    assert chapters == ["CHAPITRE 1 : PRÉSENTATION", "CHAPITRE 2 : DÉROULEMENT", "CHAPITRE 3 : BILAN"]  # capitals, accents restored
    assert sum(b["type"] == "list" for b in blocks) == 3
    assert not any("____" in (b.get("text") or "") or "……" in (b.get("text") or "") for b in blocks)


def test_shift_enter_lines_become_their_own_paragraphs(tmp_path):
    doc = Document()
    p = doc.add_paragraph()
    for k, line in enumerate(["- Absence de tableaux de bord", "", "2.2 Limites de l'existant", "Cette situation générait plusieurs problèmes pour",
                               "les équipes du terrain chaque fin de mois."]):
        run = p.add_run(line)
        if k < 4:
            run.add_break()
    path = tmp_path / "soft.docx"
    doc.save(path)
    texts = [r.text for r in from_docx(path, tmp_path / "img")]
    assert texts == ["- Absence de tableaux de bord", "2.2 Limites de l'existant", "Cette situation générait plusieurs problèmes pour les équipes du terrain chaque fin de mois."]


def test_control_characters_from_a_pdf_never_break_the_word_file(tmp_path):
    from app.doc.render import build_docx
    doc = {"blocks": [{"type": "heading", "level": 1, "text": "Chapitre\x01 un"}, {"type": "paragraph", "text": "Texte\x0b avec\x00 des codes\x1f."}], "meta": {}}
    data = build_docx(doc, tmp_path)
    assert data[:2] == b"PK"


def test_junk_rate_spots_a_bad_ocr_layer():
    from app.doc.extract import junk_rate
    assert junk_rate("Analysez ce sujet en faisant ressortir la thèse de l'auteur (2 pts) ; 1) Le domaine.") < 0.05
    assert junk_rate("~,,~ __ ~~ Comtnente~~ ut dtJ cOlTespond faisantres50rtir ~.~ .,-':;~--") > 0.5


def test_word_equations_survive_the_rebuild(tmp_path):
    from docx import Document
    from docx.oxml import parse_xml
    from app.doc.extract import from_docx
    from app.doc.render import build_docx
    import io, zipfile
    m = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'
    d = Document()
    d.add_paragraph("On considère l'équation suivante :")
    p = d.add_paragraph(); p._p.append(parse_xml(f'<m:oMathPara {m}><m:oMath><m:r><m:t>ax²+bx+c=0</m:t></m:r></m:oMath></m:oMathPara>'))
    p = d.add_paragraph("Le discriminant vaut "); p._p.append(parse_xml(f'<m:oMath {m}><m:r><m:t>Δ=b²−4ac</m:t></m:r></m:oMath>')); p.add_run(" ici.")
    d.save(tmp_path / "maths.docx")
    raws = from_docx(tmp_path / "maths.docx", tmp_path / "images")
    texts = [r.text for r in raws]
    assert texts[1].startswith("⟦eq") and texts[2].startswith("Le discriminant vaut ⟦eq") and texts[2].endswith(" ici.")
    blocks = [{"type": "paragraph", "text": t} for t in texts] + [{"type": "heading", "level": 1, "text": "Titre " + texts[1]}]
    data = build_docx({"blocks": blocks, "meta": {}}, tmp_path / "images")
    xml = zipfile.ZipFile(io.BytesIO(data)).read("word/document.xml").decode()
    assert xml.count("<m:oMath") >= 2 and "Δ=b²−4ac" in xml and "⟦" not in xml  # real equations; the title gets plain text
