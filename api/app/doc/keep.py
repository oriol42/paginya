"""Keep the user's own Word file and change only what they asked for.

`inspect` looks at a .docx and says what it already has (cover page, sommaire, page numbers and how
they are numbered, titles); `touch` writes a copy where, element by element, the cover page, the
sommaire and the page numbers are kept, added or redone. Everything else (text, styles, margins,
fonts, tables) is never touched.

Numbering follows the Cameroonian convention: the cover page is counted as "i" but shows no number,
the preliminary pages (sommaire, dédicace, remerciements, sigles, listes…) are in lowercase roman
numerals, and the body starts again at 1 on the Introduction.
"""
from __future__ import annotations

import copy
import io
import re
import unicodedata
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt

from .extract import _numbering_formats
from .render import _field

_COVER_WORDS = re.compile(
    r"(?i)\b(par\s*:|présenté|presente|matricule|année\s+(académique|scolaire|universitaire)|encadr|superviseur|"
    r"enseignant|république|republic|université|university|institut|faculté|école|ecole|mémoire|rapport|projet|"
    r"sujet|thème|soutenu|fili[èe]re|département|niveau|classe|stage|[ée]l[èe]ve|[ée]tudiant)"
)
_PAGE_FIELD = re.compile(r"(?<![A-Z])PAGE(?![A-Z])")
_OWN_ONLY = re.compile(r"^[\s\-–—/.|]*(page|p\.?)?[\s\-–—/.|]*(sur|of)?[\s\-–—/.|]*$", re.I)

# Titles that exist in a document without being written with Word's heading styles.
_TITLE_WORDS = re.compile(
    r"(?i)^(introduction(\s+g[ée]n[ée]rale)?|conclusion(\s+g[ée]n[ée]rale)?|bibliographie|r[ée]f[ée]rences?(\s+bibliographiques?)?|"
    r"webographie|annexes?|remerciements?|d[ée]dicace|[ée]pigraphe|r[ée]sum[ée]|abstract|sigles?(\s+et\s+abr[ée]viations)?|"
    r"avant[- ]propos|glossaire|liste\s+des\s+\S+|table\s+des\s+\S+|sommaire|"
    r"chapitre\s+[\divxlc]+|chapter\s+[\divxlc]+|partie\s+[\divxlc]+|"
    r"(premi[èe]re|deuxi[èe]me|troisi[èe]me|quatri[èe]me|cinqui[èe]me)\s+partie)\b"
)
_NO_TOC = re.compile(r"(?i)^(sommaire|table\s+des\s+mati[èe]res|d[ée]dicace|[ée]pigraphe|contents)\b")
_TOC_TITLE = re.compile(r"(?i)^\W*(sommaires?|table\s+des\s+mati[èe]res|contents?)\W*$")
_INTRO = re.compile(r"(?i)^\W*(introduction|intro)\b")
_ROMAN_ITEM = re.compile(r"^[IVXLC]+\s*[.\-–)]\s+\S")
_DECIMAL3 = re.compile(r"^\d+\.\d+\.\d+\.?\s+\S")
_DECIMAL2 = re.compile(r"^\d+\.\d+\.?\s+\S")
_DECIMAL1 = re.compile(r"^(\d+|[A-H])\s*[.\-–)]\s+\S")


def _paragraph_text(p) -> str:
    return "".join(t.text or "" for t in p.iter(qn("w:t")))


def _norm(text: str) -> str:
    folded = unicodedata.normalize("NFD", text)
    return "".join(c for c in folded if unicodedata.category(c) != "Mn").lower().strip()


def _has_page_break(p) -> bool:
    if p.find(".//" + qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
        return True
    if p.find(".//" + qn("w:pPr") + "/" + qn("w:pageBreakBefore")) is not None:
        return True
    return any(br.get(qn("w:type")) == "page" for br in p.iter(qn("w:br")))


def _has_page_field(element) -> bool:
    if element is None:
        return False
    for instr in element.iter(qn("w:instrText")):
        if _PAGE_FIELD.search(instr.text or ""):
            return True
    return any(_PAGE_FIELD.search(f.get(qn("w:instr")) or "") for f in element.iter(qn("w:fldSimple")))


def _section_has_numbers(section) -> bool:
    return any(_has_page_field(part._element) for part in (section.footer, section.first_page_footer, section.header, section.first_page_header))


def _style_name(paragraph) -> str:
    return (paragraph.style.name if paragraph.style is not None else "").lower()


def _is_toc_style(paragraph) -> bool:
    return _style_name(paragraph).startswith(("toc", "tm", "table des mati", "sommaire", "table of contents"))


def _is_heading(paragraph) -> bool:
    name = _style_name(paragraph)
    return name.startswith(("heading", "titre")) and not name.startswith(("titre de", "titre du"))


def _own_text(p) -> str:
    """The paragraph's own words, without those of a text box anchored to it."""
    return "".join(t.text or "" for t in p.iter(qn("w:t")) if not any(a.tag == qn("w:txbxContent") for a in t.iterancestors()))


def _flow(doc) -> list:
    """The paragraphs and tables of the body, in order: what LibreOffice lists when it reports the pages."""
    return [c for c in doc.element.body.iterchildren() if c.tag in (qn("w:p"), qn("w:tbl"))]


def page_map(doc, layout: list | None) -> dict:
    """{paragraph element: its real page} from LibreOffice's layout ([[page, text]] per paragraph or table).

    The two lists normally line up one to one. When they do not, only the paragraphs whose text is found
    again, in order, get a page: empty ones stay unknown and nothing is decided from them.
    """
    if not layout:
        return {}
    flow = _flow(doc)

    def key(text: str) -> str:
        return re.sub(r"\s+", " ", text).strip()[:40]

    pages: dict = {}
    if len(flow) == len(layout):
        for el, (page, _) in zip(flow, layout):
            if el.tag == qn("w:p") and page:
                pages[el] = int(page)
        return pages
    j = 0
    for el in flow:
        if el.tag != qn("w:p"):
            continue
        want = key(_own_text(el))
        if not want:
            continue
        for k in range(j, min(j + 12, len(layout))):
            if layout[k][0] and key(str(layout[k][1])) == want:
                pages[el] = int(layout[k][0])
                j = k + 1
                break
    return pages


def _is_blank(el) -> bool:
    """An empty line: a paragraph that only pushes what follows further down."""
    if el.tag != qn("w:p") or _paragraph_text(el).strip():
        return False
    if _has_page_break(el) or el.find(qn("w:pPr") + "/" + qn("w:numPr")) is not None:
        return False
    return not any(node.tag in (qn("w:drawing"), qn("w:pict"), qn("w:object"), qn("w:fldChar")) for node in el.iter())


def _header_text(doc) -> str:
    """Everything printed in the headers and footers (a cover's letterhead often sits there)."""
    parts = [rel.target_part for rel in doc.part.rels.values() if rel.reltype.endswith(("/header", "/footer"))]
    return "\n".join(_paragraph_text(p) for part in parts for p in part.element.iter(qn("w:p")))


def _cover_end_on_page(doc, pages: dict) -> int | None:
    """The cover is the first page as it is really laid out: few short lines, the usual words, then a new page."""
    paragraphs = list(doc.element.body.iterchildren(qn("w:p")))
    nxt = next((i for i, p in enumerate(paragraphs) if pages.get(p, 0) >= 2), None)
    if not nxt:
        return None
    while nxt < len(paragraphs) - 1 and _is_blank(paragraphs[nxt]):
        nxt += 1  # empty lines that spilled onto the second page still belong to the push
    first = paragraphs[nxt]
    lines: list[str] = []
    for child in doc.element.body.iterchildren():
        if child is first:
            break
        for p in child.iter(qn("w:p")):
            text = _own_text(p).strip()
            if text and text not in lines:
                lines.append(text)
    if len(lines) < 3 or max(len(t.split()) for t in lines) > 60 or sum(len(t.split()) for t in lines) > 300:
        return None
    score = sum(1 for t in lines + _header_text(doc).splitlines() if _COVER_WORDS.search(t))
    trailing = 0
    for p in reversed(paragraphs[:nxt]):
        if not _is_blank(p):
            break
        trailing += 1
    pushed = trailing >= 2 or any(_has_page_break(p) for p in paragraphs[:nxt])
    pushed = pushed or first.find(qn("w:pPr") + "/" + qn("w:pageBreakBefore")) is not None
    return nxt if score >= 2 and pushed else None


def _cover_end(doc, pages: dict | None = None) -> int | None:
    """Index of the first paragraph after the cover page, or None when the file has no cover page.

    With the real pages (`pages`), the first page is looked at as it prints: a cover made with empty lines
    instead of a page break is found too. Without them, only a cover closed by a page break is.
    """
    if pages:
        return _cover_end_on_page(doc, pages)
    body = list(doc.element.body.iterchildren(qn("w:p")))[:80]
    lines = 0
    score = 0
    for i, p in enumerate(body):
        text = _paragraph_text(p).strip()
        if text:
            lines += 1
            if _COVER_WORDS.search(text):
                score += 1
        if _has_page_break(p):
            if lines >= 3 and score >= 2 and i < 70:
                return i + 1
            return None
    return None


# --- Titles written without Word's heading styles ---------------------------------------------

def _runs_bold(paragraph) -> bool:
    runs = [r for r in paragraph.runs if (r.text or "").strip()]
    if not runs:
        return False
    style_bold = bool(paragraph.style is not None and paragraph.style.font is not None and paragraph.style.font.bold)
    return all(r.bold if r.bold is not None else style_bold for r in runs)


def _runs_size(paragraph) -> float:
    sizes = [r.font.size.pt for r in paragraph.runs if r.font.size is not None and (r.text or "").strip()]
    if sizes:
        return max(sizes)
    style = paragraph.style
    return style.font.size.pt if style is not None and style.font is not None and style.font.size is not None else 0.0


def _outline_level(paragraph) -> int | None:
    ppr = paragraph._p.pPr
    node = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
    if node is not None:
        try:
            level = int(node.get(qn("w:val")))
        except (TypeError, ValueError):
            return None
        return level + 1 if level < 9 else None
    return None


def _num_of(paragraph) -> tuple[str, str] | None:
    """(numId, ilvl) when the paragraph is numbered by Word itself ("I.", "1.", "a)" typed by nobody)."""
    node = paragraph._p.find(qn("w:pPr") + "/" + qn("w:numPr"))
    num_id = node.find(qn("w:numId")) if node is not None else None
    if num_id is None or num_id.get(qn("w:val")) in (None, "0"):
        return None
    level = node.find(qn("w:ilvl"))
    return num_id.get(qn("w:val")), (level.get(qn("w:val")) if level is not None else "0")


def find_titles(doc, after: int = 0) -> list[tuple[object, int]]:
    """[(paragraph, level)] for the titles of the document, however they are written.

    A paragraph is a title when it uses a heading style, has an outline level, starts like a known
    section ("INTRODUCTION", "CHAPITRE I", "I. …", "1.1 …") or is a short line in bold/capitals/centred.
    The sommaire itself and the cover page (the first `after` paragraphs) are never titles.
    """
    paragraphs = doc.paragraphs
    found: list[tuple[object, int]] = []
    formats = _numbering_formats(doc)
    numbered: dict[tuple[str, str], int] = {}  # Word's own numbering: each list/level met is one level deeper
    for i, p in enumerate(paragraphs):
        if i < after or _is_toc_style(p):
            continue
        text = (p.text or "").strip()
        if not text or len(text) > 140:
            continue
        if _is_heading(p):
            digits = re.search(r"(\d)\s*$", _style_name(p))
            found.append((p, min(int(digits.group(1)), 4) if digits else 1))
            continue
        outline = _outline_level(p)
        if outline:
            found.append((p, min(outline, 4)))
            continue
        if text.endswith((".", ",", ";")) and not _ROMAN_ITEM.match(text) and not _DECIMAL1.match(text) and not _DECIMAL2.match(text):
            continue
        words = len(text.split())
        bold = _runs_bold(p)
        centered = p.alignment == WD_ALIGN_PARAGRAPH.CENTER
        big = _runs_size(p) >= 13
        caps = text.upper() == text and sum(c.isalpha() for c in text) >= 4
        num = _num_of(p)
        known = _TITLE_WORDS.match(text)
        if known and not re.search(r"(?i)chap|partie", known.group()) and re.match(r"\s*:\s*\S", text[known.end():]):
            continue  # "Conclusion : valide" is a sentence of the text, not the Conclusion
        if num and formats.get(num, "bullet") not in ("bullet", "none") and bold and words <= 16:
            found.append((p, min(numbered.setdefault(num, len(numbered) + 1), 3)))
        elif known and words <= 14:
            found.append((p, 1))
        elif _DECIMAL3.match(text) and bold and words <= 16:
            found.append((p, 3))
        elif _DECIMAL2.match(text) and (bold or big) and words <= 16:
            found.append((p, 2))
        elif _ROMAN_ITEM.match(text) and (bold or caps or big) and words <= 16:
            found.append((p, 1 if caps or big else 2))
        elif _DECIMAL1.match(text) and (bold or big) and words <= 16:
            found.append((p, 2))
        elif caps and (bold or big or centered) and words <= 12:
            found.append((p, 1))
    return found


def _titles_after_cover(doc, cover_end: int | None) -> list[tuple[object, int]]:
    """Titles of the body (the cover page and the sommaire are left out)."""
    return [(p, lv) for p, lv in find_titles(doc, after=cover_end or 0) if not _TOC_TITLE.match((p.text or "").strip())]


# --- What the file already has ------------------------------------------------------------------

def _numbering_of(doc) -> dict:
    """How the pages are numbered now: {"scheme": none|arabic|roman|roman_arabic, "on_cover": bool}."""
    formats: list[str] = []
    numbered = False
    for section in doc.sections:
        has = _section_has_numbers(section)
        numbered = numbered or has
        pg = section._sectPr.find(qn("w:pgNumType"))
        formats.append((pg.get(qn("w:fmt")) if pg is not None else None) or "decimal")
    if not numbered:
        return {"scheme": "none", "on_cover": False}
    has_roman = any("Roman" in f for f in formats)
    has_arabic = any("Roman" not in f for f in formats)
    scheme = "roman_arabic" if has_roman and has_arabic else "roman" if has_roman else "arabic"
    first = doc.sections[0]
    on_cover = _has_page_field(first.footer._element) and not first.different_first_page_header_footer
    on_cover = on_cover or (_has_page_field(first.first_page_footer._element) and first.different_first_page_header_footer)
    return {"scheme": scheme, "on_cover": bool(on_cover)}


def _own_layout(doc) -> bool:
    """Things placed by hand that only exist in this file: text boxes, shapes, a letterhead, a page border."""
    body = doc.element.body
    if any(node.tag in (qn("w:txbxContent"), qn("w:pict")) for node in body.iter()):
        return True
    if any(True for _ in body.iter(qn("w:pgBorders"))):
        return True
    for rel in doc.part.rels.values():
        if rel.reltype.endswith(("/header", "/footer")):
            element = rel.target_part.element
            if _paragraph_text(element).strip() or any(True for _ in element.iter(qn("w:drawing"))):
                return True
    return False


def _empty_toc_title(doc, cover_end: int | None):
    """The "SOMMAIRE" written above a page left empty (to be filled by hand), or None."""
    paragraphs = list(doc.element.body.iterchildren(qn("w:p")))
    for p in paragraphs[cover_end or 0:]:
        if _TOC_TITLE.match(_own_text(p).strip()):
            nxt = p.getnext()
            return None if nxt is not None and _paragraph_text(nxt).startswith("[[TOC") else p
    return None


def inspect(path: Path, layout: list | None = None) -> dict:
    """What the Word file already has: cover, sommaire, page numbers (and their scheme), titles, sections.

    `layout` is LibreOffice's real pagination (office.layout): with it the cover is the page that prints first.
    """
    doc = Document(str(path))
    body = doc.element.body
    toc = any(re.match(r"\s*TOC\b", i.text or "") for i in body.iter(qn("w:instrText")))
    toc = toc or any((f.get(qn("w:instr")) or "").strip().startswith("TOC") for f in body.iter(qn("w:fldSimple")))
    if not toc:
        toc = any(_is_toc_style(p) for p in doc.paragraphs)
    cover_end = _cover_end(doc, page_map(doc, layout))
    titles = _titles_after_cover(doc, cover_end)
    numbering = _numbering_of(doc)
    return {
        "cover": cover_end is not None,
        "toc": bool(toc),
        "page_numbers": any(_section_has_numbers(s) for s in doc.sections),
        "sections": len(doc.sections),
        "headings": sum(1 for p in doc.paragraphs if _is_heading(p)),
        "titles": len(titles),
        "numbering": numbering["scheme"],
        "numbers_on_cover": numbering["on_cover"],
        "intro": any(_INTRO.match((p.text or "").strip()) for p, lv in titles if lv == 1),
        "own_layout": _own_layout(doc),
        "toc_title": not toc and _empty_toc_title(doc, cover_end) is not None,
    }


def header_lines(path: Path) -> list[str]:
    """The text of the first-page header (the printed letterhead of many covers), line by line."""
    doc = Document(str(path))
    section = doc.sections[0]
    parts = [section.first_page_header, section.header] if section.different_first_page_header_footer else [section.header]
    lines: list[str] = []
    for part in parts:
        for p in part.paragraphs:
            lines.extend(t.strip() for t in re.split(r"\n", p.text or "") if t.strip())
        for table in part.tables:
            for row in table.rows:
                for cell in row.cells:
                    lines.extend(t.strip() for t in re.split(r"\n", cell.text or "") if t.strip())
        if lines:
            break
    return lines


# --- Editing helpers ----------------------------------------------------------------------------

def _strip_number_fields(element) -> None:
    """Removes the PAGE fields (and NUMPAGES) of a header/footer; a paragraph left with only "Page … sur" goes too."""
    for p in list(element.iter(qn("w:p"))):
        changed = False
        for fld in list(p.iter(qn("w:fldSimple"))):
            if re.search(r"\b(PAGE|NUMPAGES)\b", fld.get(qn("w:instr")) or ""):
                fld.getparent().remove(fld)
                changed = True
        depth, drop, runs = 0, False, []
        for r in list(p.iter(qn("w:r"))):
            fc = r.find(qn("w:fldChar"))
            if fc is not None:
                kind = fc.get(qn("w:fldCharType"))
                if kind == "begin":
                    depth += 1
                    drop = False
                    runs = [r]
                    continue
                if kind == "end" and depth:
                    runs.append(r)
                    depth -= 1
                    if drop:
                        for rr in runs:
                            if rr.getparent() is not None:
                                rr.getparent().remove(rr)
                        changed = True
                    runs = []
                    continue
            if depth:
                runs.append(r)
                instr = r.find(qn("w:instrText"))
                if instr is not None and re.search(r"\b(PAGE|NUMPAGES)\b", instr.text or ""):
                    drop = True
        if changed and _OWN_ONLY.match(_paragraph_text(p)):
            parent = p.getparent()
            if parent is not None and len([c for c in parent.iterchildren(qn("w:p"))]) > 1:
                parent.remove(p)
            else:
                for child in list(p):
                    if child.tag != qn("w:pPr"):
                        p.remove(child)


def _drop_refs(sectpr, kinds: tuple[str, ...], tags=("w:headerReference", "w:footerReference")) -> None:
    for tag in tags:
        for ref in sectpr.findall(qn(tag)):
            if ref.get(qn("w:type")) in kinds:
                sectpr.remove(ref)


def _set_pgnum(sectpr, fmt: str, start: int | None) -> None:
    pg = sectpr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        # schema order: pgNumType comes after pgBorders/lnNumType and before cols
        anchor = next((sectpr.find(qn(t)) for t in ("w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg", "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid") if sectpr.find(qn(t)) is not None), None)
        if anchor is not None:
            anchor.addprevious(pg)
        else:
            sectpr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is None:
        pg.attrib.pop(qn("w:start"), None)
    else:
        pg.set(qn("w:start"), str(start))


def _remove_break_at_end(p) -> None:
    """A page break that ends the paragraph is replaced by the section break we are adding."""
    runs = [r for r in p.iter(qn("w:r"))]
    for r in reversed(runs):
        brs = [b for b in r.findall(qn("w:br")) if b.get(qn("w:type")) == "page"]
        if brs:
            for b in brs:
                r.remove(b)
            if not [c for c in r if c.tag != qn("w:rPr")]:
                r.getparent().remove(r)
            return
        if (r.find(qn("w:t")) is not None and (r.find(qn("w:t")).text or "").strip()):
            return


def _end_section_after(doc, element, final_sectpr) -> object:
    """Makes a section end at `element` (a paragraph, or whatever precedes the break). Returns its sectPr.

    The new section inherits the layout of the last one (a copy of its sectPr). An existing section
    break at that place is reused.
    """
    if element.tag == qn("w:p"):
        existing = element.find(qn("w:pPr") + "/" + qn("w:sectPr"))
        if existing is not None:
            return existing
        _remove_break_at_end(element)
        holder = element
    else:
        holder = OxmlElement("w:p")
        element.addnext(holder)
    ppr = holder.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr")
        holder.insert(0, ppr)
    sect = copy.deepcopy(final_sectpr)
    ppr.append(sect)  # sectPr is the last child of pPr
    kind = sect.find(qn("w:type"))
    if kind is None:
        kind = OxmlElement("w:type")
        refs = len(sect.findall(qn("w:headerReference"))) + len(sect.findall(qn("w:footerReference")))
        sect.insert(refs, kind)
    kind.set(qn("w:val"), "nextPage")
    return sect


def _section_ends(doc) -> list[int]:
    """For each section, the index (among top-level paragraphs) of the paragraph that ends it."""
    ends: list[int] = []
    pi = -1
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            pi += 1
            if child.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None:
                ends.append(pi)
    ends.append(10**9)
    return ends


def _own_footer(section, paragraphs_to_keep: list, number: bool) -> None:
    """Gives the section a footer of its own: the text it had (without page numbers), then the page number."""
    _drop_refs(section._sectPr, ("default",), ("w:footerReference",))
    footer = section.footer
    footer.is_linked_to_previous = False
    body = footer._element
    for p in list(body.iterchildren(qn("w:p"))):
        body.remove(p)
    for p in paragraphs_to_keep:
        body.append(copy.deepcopy(p))
    if number:
        para = footer.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _field(para, "PAGE")
        for run in para.runs:
            run.font.size = Pt(10)
    elif not body.findall(qn("w:p")):
        footer.add_paragraph()


def _kept_footer_paragraphs(section) -> list:
    """The paragraphs of the section's footer that are not page numbers (their own text is never lost)."""
    scratch = copy.deepcopy(section.footer._element)
    _strip_number_fields(scratch)
    return [p for p in scratch.iterchildren(qn("w:p")) if _paragraph_text(p).strip() or p.findall(".//" + qn("w:drawing"))]


def _strip_first_page(sectpr) -> None:
    """A section that is not the cover has no 'different first page' (that was the cover's letterhead)."""
    for node in sectpr.findall(qn("w:titlePg")):
        sectpr.remove(node)
    _drop_refs(sectpr, ("first",))


# --- Cover page ---------------------------------------------------------------------------------

def _replace_cover(doc, cover_end: int | None, png: bytes) -> None:
    """Paginya's cover (a full-page picture of its own section) in place of the old one, or at the very start."""
    body = doc.element.body
    final = body.find(qn("w:sectPr"))
    if cover_end:
        paragraphs = list(body.iterchildren(qn("w:p")))
        last = paragraphs[cover_end - 1]
        for child in list(body.iterchildren()):
            if child.tag == qn("w:sectPr"):
                break
            stop = child is last
            body.remove(child)
            if stop:
                break
    cover_p = OxmlElement("w:p")
    body.insert(0, cover_p)
    sect = copy.deepcopy(final)
    for tag in ("w:pgBorders", "w:titlePg"):
        for node in sect.findall(qn(tag)):
            sect.remove(node)
    for ref in list(sect.findall(qn("w:headerReference"))) + list(sect.findall(qn("w:footerReference"))):
        sect.remove(ref)
    pg_mar = sect.find(qn("w:pgMar"))
    if pg_mar is None:
        pg_mar = OxmlElement("w:pgMar")
        sect.append(pg_mar)
    for attr in ("top", "right", "bottom", "left", "header", "footer", "gutter"):
        pg_mar.set(qn(f"w:{attr}"), "0")
    kind = sect.find(qn("w:type"))
    if kind is None:
        kind = OxmlElement("w:type")
        sect.insert(0, kind)
    kind.set(qn("w:val"), "nextPage")
    ppr = OxmlElement("w:pPr")
    ppr.append(sect)
    cover_p.append(ppr)
    from docx.text.paragraph import Paragraph

    para = Paragraph(cover_p, doc._body)
    para.paragraph_format.space_after = para.paragraph_format.space_before = Pt(0)
    para.add_run().add_picture(io.BytesIO(png), width=Mm(209.5))
    # the cover section is the first one: the body must not inherit its 'different first page'
    _strip_first_page(final)


# --- Real page breaks ---------------------------------------------------------------------------

def _starts_page(p) -> bool:
    if p.find(qn("w:pPr") + "/" + qn("w:pageBreakBefore")) is not None:
        return True
    previous = p.getprevious()
    return previous is not None and previous.tag == qn("w:p") and _has_page_break(previous)


def _page_break_before(doc, p) -> None:
    from docx.text.paragraph import Paragraph

    Paragraph(p, doc._body).paragraph_format.page_break_before = True


def _real_page_breaks(doc, pages: dict, after_cover=None) -> int:
    """Empty lines typed until the next page are replaced by a page break. Returns how many places changed.

    Only a run of empty lines that really crosses a page is touched (three lines or more, or the end of the
    cover): spacing inside a page stays as it is.
    """
    changed = 0
    run: list = []
    for el in _flow(doc):
        if _is_blank(el):
            run.append(el)
            continue
        if run and el.tag == qn("w:p") and pages.get(el) and pages.get(run[0]) and pages[el] > pages[run[0]]:
            if len(run) >= 3 or el is after_cover:
                for blank in run:
                    blank.getparent().remove(blank)
                if not _starts_page(el):
                    _page_break_before(doc, el)
                changed += 1
        run = []
    return changed


# --- Sommaire -----------------------------------------------------------------------------------

def _toc_field_paragraphs(doc) -> list | None:
    """The paragraphs (and wrapper, if the sommaire sits in a content control) of an existing Word sommaire."""
    body = doc.element.body
    for sdt in body.iterchildren(qn("w:sdt")):
        gallery = sdt.find(".//" + qn("w:docPartGallery"))
        if gallery is not None and "table of contents" in (gallery.get(qn("w:val")) or "").lower():
            return [sdt]
    items = list(body.iterchildren())
    start = end = None
    depth = 0
    for i, el in enumerate(items):
        if el.tag != qn("w:p"):
            continue
        for node in el.iter():
            if node.tag == qn("w:instrText") and start is None and re.match(r"\s*TOC\b", node.text or ""):
                start = i
            if node.tag == qn("w:fldSimple") and start is None and (node.get(qn("w:instr")) or "").strip().startswith("TOC"):
                start = end = i
            if start is not None and end is None and node.tag == qn("w:fldChar"):
                kind = node.get(qn("w:fldCharType"))
                if kind == "begin":
                    depth += 1
                elif kind == "end":
                    depth -= 1
                    if depth <= 0:
                        end = i
        if start is not None and end is not None:
            break
    if start is None:
        # no field: a run of paragraphs in the usual sommaire styles
        run = [el for el in items if el.tag == qn("w:p") and re.match(r"(?i)(toc|tm|table des mati|sommaire)", _style_of(doc, el))]
        return run or None
    # the field may begin in the paragraph of its title; paragraphs up to its end belong to it
    return [el for el in items[start : (end if end is not None else start) + 1]]


def _style_of(doc, p_el) -> str:
    from docx.text.paragraph import Paragraph

    try:
        return _style_name(Paragraph(p_el, doc._body))
    except Exception:
        return ""


def _mark_outline(titles: list[tuple[object, int]]) -> int:
    """Gives the detected titles an outline level (invisible), so the sommaire can list them."""
    done = 0
    for p, level in titles:
        if _is_heading(p) or _outline_level(p) or _NO_TOC.match((p.text or "").strip()):
            continue
        ppr = p._p.get_or_add_pPr()
        node = OxmlElement("w:outlineLvl")
        node.set(qn("w:val"), str(min(level, 3) - 1))
        # outlineLvl follows the paragraph properties listed before it in the schema; appending is accepted by Word and LibreOffice
        ppr.append(node)
        done += 1
    return done


def _insert_marker_paragraphs(anchor, with_title: bool, levels: int) -> None:
    doc_el = anchor
    if with_title:
        title = OxmlElement("w:p")
        ppr = OxmlElement("w:pPr")
        jc = OxmlElement("w:jc")
        jc.set(qn("w:val"), "center")
        ppr.append(jc)
        title.append(ppr)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        bold = OxmlElement("w:b")
        size = OxmlElement("w:sz")
        size.set(qn("w:val"), "32")
        rpr.append(bold)
        rpr.append(size)
        run.append(rpr)
        text = OxmlElement("w:t")
        text.text = "Sommaire"
        run.append(text)
        title.append(run)
        doc_el.addprevious(title)
    marker = OxmlElement("w:p")
    run = OxmlElement("w:r")
    text = OxmlElement("w:t")
    text.text = f"[[TOC:{levels}]]"
    run.append(text)
    marker.append(run)
    doc_el.addprevious(marker)


# --- The one entry point ------------------------------------------------------------------------

def touch(
    src: Path,
    dst: Path,
    page_numbers: bool = False,
    toc: bool | str = False,
    *,
    numbers: str = "",
    cover: str = "keep",
    cover_png: bytes | None = None,
    layout: list | None = None,
) -> list[str]:
    """Writes a copy of `src` where the cover, sommaire and page numbers are kept, added or redone.

    numbers: "keep" | "add" (only if there are none) | "redo" (Cameroonian scheme) | "none"
    toc:     "keep" | "add" | "redo" | "none"        (True means "add")
    cover:   "keep" | "redo" | "add"  (needs `cover_png`, Paginya's cover as a picture)
    page_numbers=True is the older spelling of numbers="add".
    layout:  LibreOffice's real pagination of `src` (office.layout); pages held apart by empty lines then
             get a real page break, so that adding a sommaire does not push everything out of place.
    Returns what was done, in words.
    """
    numbers = numbers or ("add" if page_numbers else "keep")
    toc = "add" if toc is True else "keep" if toc is False else toc
    doc = Document(str(src))
    done: list[str] = []
    pages = page_map(doc, layout)
    cover_end = _cover_end(doc, pages)
    if pages:
        paragraphs = list(doc.element.body.iterchildren(qn("w:p")))
        after_cover = paragraphs[cover_end] if cover_end else None
        _real_page_breaks(doc, pages, after_cover)
        if after_cover is not None:
            cover_end = list(doc.element.body.iterchildren(qn("w:p"))).index(after_cover)

    if cover in ("redo", "add") and cover_png:
        if cover == "add" and cover_end is not None:
            pass  # there is one already: "add" never replaces
        else:
            _replace_cover(doc, cover_end, cover_png)
            cover_end = 1
            done.append("page de garde")
    if cover_end is not None and cover_end > len(list(doc.element.body.iterchildren(qn("w:p")))):
        cover_end = None

    titles = _titles_after_cover(doc, cover_end)

    had_toc = bool(_toc_field_paragraphs(doc))
    if toc == "redo" and had_toc and titles:
        group = _toc_field_paragraphs(doc) or []
        anchor = group[0]
        has_title = False
        prev = anchor.getprevious()
        if prev is not None and prev.tag == qn("w:p") and _TOC_TITLE.match(_paragraph_text(prev).strip()):
            has_title = True
        _insert_marker_paragraphs(anchor, with_title=not has_title, levels=2)
        for el in group:
            if el.getparent() is not None:
                el.getparent().remove(el)
        _mark_outline(titles)
        done.append("sommaire refait")
    elif toc in ("add", "redo") and not had_toc and titles and _empty_toc_title(doc, cover_end) is not None:
        # the page is there, with its title: the sommaire goes under it, and what follows starts a new page
        _mark_outline(titles)
        holder = _empty_toc_title(doc, cover_end)
        while holder.getnext() is not None and _is_blank(holder.getnext()):
            holder.getparent().remove(holder.getnext())
        following = holder.getnext()
        marker = OxmlElement("w:p")
        holder.addnext(marker)
        _insert_marker_paragraphs(marker, with_title=False, levels=2)
        marker.getparent().remove(marker)
        if following is not None and following.tag == qn("w:p") and not _starts_page(following):
            _page_break_before(doc, following)
        done.append("sommaire")
    elif toc in ("add", "redo") and not had_toc and titles and not any(p.text.startswith("[[TOC") for p in doc.paragraphs):
        _mark_outline(titles)
        first = next((p for p, _ in titles if p._p.getprevious() is not None), titles[0][0])
        _insert_marker_paragraphs(first._p, with_title=True, levels=2)
        br = OxmlElement("w:p")
        run = OxmlElement("w:r")
        brk = OxmlElement("w:br")
        brk.set(qn("w:type"), "page")
        run.append(brk)
        br.append(run)
        first._p.addprevious(br)
        done.append("sommaire")

    if numbers in ("add", "redo"):
        if _apply_numbers(doc, cover_end, redo=numbers == "redo"):
            done.append("numéros de page")

    dst.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(dst))
    return done


def _apply_numbers(doc, cover_end: int | None, redo: bool) -> bool:
    """Cameroonian numbering: cover counted but blank, preliminary pages i, ii…, body 1, 2… from the Introduction."""
    if not redo and any(_section_has_numbers(s) for s in doc.sections):
        return False
    body = doc.element.body
    final = body.find(qn("w:sectPr"))
    if final is None:
        return False
    titles = _titles_after_cover(doc, cover_end)
    paragraphs = list(body.iterchildren(qn("w:p")))
    intro_el = next((p._p for p, lv in titles if lv == 1 and _INTRO.match((p.text or "").strip())), None)
    intro_index = paragraphs.index(intro_el) if intro_el is not None and intro_el in paragraphs else None
    # prelim pages exist when something sits between the cover and the Introduction
    start_after_cover = cover_end or 0
    has_prelim = intro_index is not None and intro_index > start_after_cover and any(
        _paragraph_text(p).strip() for p in paragraphs[start_after_cover:intro_index]
    )

    created: list = []
    if cover_end:
        before = paragraphs[cover_end - 1]
        if before.find(qn("w:pPr") + "/" + qn("w:sectPr")) is None:
            created.append(_end_section_after(doc, before, final))
    if has_prelim and intro_el is not None:
        previous = intro_el.getprevious()
        if previous is not None and not (previous.tag == qn("w:p") and previous.find(qn("w:pPr") + "/" + qn("w:sectPr")) is not None):
            created.append(_end_section_after(doc, previous, final))
    if created:
        # sections cut out of a file that had a cover on its first page: only the first one keeps that
        for sect in created[1:]:
            _strip_first_page(sect)
        _strip_first_page(final)

    paragraphs = list(body.iterchildren(qn("w:p")))
    ends = _section_ends(doc)
    cover_last = (cover_end - 1) if cover_end else -1
    intro_pos = paragraphs.index(intro_el) if intro_el is not None and intro_el in paragraphs else None
    sections = list(doc.sections)
    roles: list[str] = []
    for i, end in enumerate(ends[: len(sections)]):
        if cover_end and end <= cover_last:
            roles.append("cover")
        elif has_prelim and intro_pos is not None and end < intro_pos:
            roles.append("prelim")
        else:
            roles.append("body")
    if "body" not in roles:
        roles[-1] = "body"

    if redo:  # every number of the old scheme goes first (headers too: "Page 3" at the top)
        for section in sections:
            for part in (section.header, section.first_page_header, section.even_page_header):
                if not part.is_linked_to_previous:
                    _strip_number_fields(part._element)

    first_body = True
    for section, role in zip(sections, roles):
        keep_text = _kept_footer_paragraphs(section)
        if role == "cover":
            _set_pgnum(section._sectPr, "lowerRoman", 1)
            _own_footer(section, [] if redo else keep_text, number=False)
            first = section.first_page_footer
            if not first.is_linked_to_previous:
                _strip_number_fields(first._element)
        elif role == "prelim":
            _set_pgnum(section._sectPr, "lowerRoman", None if cover_end else 1)
            _own_footer(section, keep_text, number=True)
        else:
            _set_pgnum(section._sectPr, "decimal", 1 if first_body else None)
            first_body = False
            _own_footer(section, keep_text, number=True)
        if role != "cover" and (created or cover_end) and (len(sections) > 1):
            _strip_first_page(section._sectPr)
    return True


def inspect_toc(doc) -> bool:
    body = doc.element.body
    return any(re.match(r"\s*TOC\b", i.text or "") for i in body.iter(qn("w:instrText")))
