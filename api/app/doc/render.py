"""Blocks + style -> .docx (python-docx).

Tables of contents and lists of tables/figures are left as [[MARKERS]]:
LibreOffice turns them into real indexes with page numbers (office.py).
"""
from __future__ import annotations

import base64
import io
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor

from .clean import typo_fr
from PIL import Image

from .. import render as cover_render

PRELIM_SPECIALS = {"dedicace", "epigraphe", "remerciements", "avant_propos", "sigles", "resume", "abstract"}
NO_TOC_SPECIALS = {"dedicace", "epigraphe"}  # never listed in the table of contents

THEMES: dict[str, dict] = {
    "academique": {
        "font": "Times New Roman", "size": 12, "line": 1.5, "justify": True, "indent": 1.25,
        "heading_color": "000000", "accent": "000000", "caps_h1": True, "center_h1": True,
        "sizes": (16, 14, 12, 12), "table_head": "E7E6E6", "table_head_text": "000000",
        "page_number": "center", "rule_h1": False,
    },
    "moderne": {
        "font": "Calibri", "size": 11, "line": 1.15, "justify": True, "indent": 0,
        "heading_color": None, "accent": None, "caps_h1": False, "center_h1": False,
        "sizes": (20, 15, 12.5, 11.5), "table_head": None, "table_head_text": "FFFFFF",
        "page_number": "right", "rule_h1": True,
    },
    "simple": {
        "font": "Arial", "size": 11, "line": 1.15, "justify": False, "indent": 0,
        "heading_color": "111111", "accent": "444444", "caps_h1": False, "center_h1": False,
        "sizes": (16, 13, 11.5, 11), "table_head": "F2F2F2", "table_head_text": "000000",
        "page_number": "center", "rule_h1": False,
    },
    # Times, like the academic norm, but titles in colour with a fine rule.
    "universitaire": {
        "font": "Times New Roman", "size": 12, "line": 1.5, "justify": True, "indent": 1.25,
        "heading_color": None, "accent": None, "caps_h1": True, "center_h1": True,
        "sizes": (16, 14, 12, 12), "table_head": None, "table_head_text": "FFFFFF",
        "page_number": "center", "rule_h1": True,
    },
    # Cambria (Caladea on the server), small caps titles, airy and refined.
    "elegant": {
        "font": "Cambria", "size": 11.5, "line": 1.3, "justify": True, "indent": 0.8,
        "heading_color": None, "accent": None, "caps_h1": False, "center_h1": True,
        "sizes": (20, 14, 12, 11.5), "table_head": "F3F4F6", "table_head_text": "111111",
        "page_number": "center", "rule_h1": True, "small_caps": True,
    },
    # Business report: level-1 titles in a coloured band with white text.
    "corporate": {
        "font": "Calibri", "size": 11, "line": 1.15, "justify": True, "indent": 0,
        "heading_color": None, "accent": None, "caps_h1": True, "center_h1": False,
        "sizes": (15, 14, 12, 11), "table_head": None, "table_head_text": "FFFFFF",
        "page_number": "right", "rule_h1": False, "band_h1": True,
    },
}

DEFAULT_STYLE = {
    "theme": "academique", "font": None, "size": None, "line": None, "justify": None,
    "margins": [2.5, 2.5, 3.0, 2.5],  # top, bottom, left, right (cm)
    "color": "#0E9F6E",
}
DEFAULT_OPTIONS = {"toc": False, "toc_end": False, "lists": False, "cover": False, "page_numbers": True, "chapter_pages": False, "letterhead": False}

# Kinds of document the user can pick (detected first, always changeable). Each one only adds the
# pages its kind really has: a course or a letter never gets a table of contents by itself.
KINDS = {
    "memoire": {"toc": True, "toc_end": True, "lists": True, "cover": True, "page_numbers": True, "chapter_pages": True},
    "rapport_stage": {"toc": True, "toc_end": True, "lists": True, "cover": True, "page_numbers": True, "chapter_pages": True},
    "rapport": {"toc": True, "toc_end": False, "lists": True, "cover": True, "page_numbers": True, "chapter_pages": True},
    "expose": {"toc": False, "toc_end": False, "lists": False, "cover": True, "page_numbers": True, "chapter_pages": False},
    "cours": {"toc": False, "toc_end": False, "lists": False, "cover": False, "page_numbers": True, "chapter_pages": False},
    "administratif": {"toc": False, "toc_end": False, "lists": False, "cover": False, "page_numbers": False, "chapter_pages": False},
    "document": {"toc": False, "toc_end": False, "lists": False, "cover": False, "page_numbers": True, "chapter_pages": False},
}


def options_for(kind: str, words: int = 0) -> dict:
    opts = dict(KINDS.get(kind, KINDS["document"]))
    if kind == "document" and words < 600:
        opts["page_numbers"] = False  # a one-page text needs no "1"
    return opts


def resolved_style(style: dict | None) -> dict:
    s = {**DEFAULT_STYLE, **(style or {})}
    theme = THEMES.get(s["theme"], THEMES["academique"])
    color = (s.get("color") or "#0E9F6E").lstrip("#").upper()
    return {
        "theme": s["theme"] if s["theme"] in THEMES else "academique",
        "font": s["font"] or theme["font"],
        "size": float(s["size"] or theme["size"]),
        "line": float(s["line"] or theme["line"]),
        "justify": theme["justify"] if s["justify"] is None else bool(s["justify"]),
        "indent": theme["indent"],
        "margins": [float(x) for x in (s.get("margins") or DEFAULT_STYLE["margins"])][:4],
        "heading_color": theme["heading_color"] or color,
        "accent": theme["accent"] or color,
        "caps_h1": theme["caps_h1"],
        "center_h1": theme["center_h1"],
        "sizes": theme["sizes"],
        "table_head": theme["table_head"] or color,
        "table_head_text": theme["table_head_text"],
        "page_number": theme["page_number"],
        "rule_h1": theme["rule_h1"],
        "small_caps": theme.get("small_caps", False),
        "band_h1": theme.get("band_h1", False),
    }


# --- low level OOXML helpers ------------------------------------------------

def _set_font(rpr_owner, name: str) -> None:
    """Force a font on a style/run, removing theme fonts that would override it."""
    rpr = rpr_owner.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        fonts.attrib.pop(qn(attr), None)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attr), name)


def _set_color(font, hex6: str) -> None:
    font.color.rgb = RGBColor.from_string(hex6)
    color = font.element.rPr.find(qn("w:color")) if font.element.rPr is not None else None
    if color is not None:
        for attr in ("w:themeColor", "w:themeShade", "w:themeTint"):
            color.attrib.pop(qn(attr), None)


def _field(paragraph, instr: str, placeholder: str = "1") -> None:
    run = paragraph.add_run()
    for kind in ("begin", "instr", "separate", "text", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instr} "
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = placeholder
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def _shade(cell, hex6: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    tcpr.append(shd)


def _bottom_rule(paragraph, hex6: str, size: int = 12) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", str(size)), ("w:space", "4"), ("w:color", hex6)):
        bottom.set(qn(k), v)
    bdr.append(bottom)
    ppr.append(bdr)


def _top_rule(paragraph, hex6: str) -> None:
    bdr = paragraph._p.get_or_add_pPr().find(qn("w:pBdr"))
    if bdr is None:
        bdr = OxmlElement("w:pBdr")
        paragraph._p.get_or_add_pPr().append(bdr)
    top = OxmlElement("w:top")
    for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "6"), ("w:color", hex6)):
        top.set(qn(k), v)
    bdr.insert(0, top)


def _band(paragraph, hex6: str) -> None:
    """Coloured background behind a paragraph (corporate titles)."""
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    ppr.append(shd)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), "113")
    ind.set(qn("w:right"), "113")
    ppr.append(ind)


_PPR_AFTER_SHD = ("tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct", "topLinePunct", "autoSpaceDE",
                  "autoSpaceDN", "bidi", "adjustRightInd", "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
                  "suppressOverlap", "jc", "textDirection", "textAlignment", "textboxTightWrap", "outlineLvl", "divId", "cnfStyle",
                  "rPr", "sectPr", "pPrChange")


def _shade_para(paragraph, hex6: str) -> None:
    """Paragraph background, inserted where the OOXML schema wants it (Word rejects misplaced elements)."""
    ppr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    after = next((c for c in ppr if c.tag.split("}")[-1] in _PPR_AFTER_SHD), None)
    if after is None:
        ppr.append(shd)
    else:
        after.addprevious(shd)


def _no_borders(table) -> None:
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        borders.append(el)
    tblpr.append(borders)


def _full_width(table) -> None:
    tblpr = table._tbl.tblPr
    width = tblpr.find(qn("w:tblW"))
    if width is None:
        width = OxmlElement("w:tblW")
        tblpr.append(width)
    width.set(qn("w:type"), "pct")
    width.set(qn("w:w"), "5000")


def _page_numbers(section, fmt: str, start: int | None, align: str, font: str, color: str) -> None:
    sectpr = section._sectPr
    pg = sectpr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectpr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    if start is not None:
        pg.set(qn("w:start"), str(start))
    section.footer.is_linked_to_previous = False
    p = section.footer.paragraphs[0]
    p.text = ""
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT if align == "right" else WD_ALIGN_PARAGRAPH.CENTER
    _field(p, "PAGE")
    for run in p.runs:
        _set_font(run._r, font)
        run.font.size = Pt(10)
        _set_color(run.font, color)


LINK = re.compile(r"(https?://[^\s)»]+[^\s.,;:)»]|www\.[^\s)»]+[^\s.,;:)»]|[\w.+-]+@[\w-]+\.[\w.-]*\w)")


INLINE = re.compile(r"(\*\*\*|\*\*|__)(?=\S)(.+?)(?<=\S)\1|(?<![\w*])\*(?=[^\s*])(.+?)(?<=[^\s*])\*(?![\w*])|`([^`\n]+)`")


def inline_spans(text: str) -> list[tuple[str, str]]:
    """Markdown-style emphasis → [(text, flags)] with flags among "b" (bold), "i" (italic), "c" (code)."""
    out, pos = [], 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            out.append((text[pos:m.start()], ""))
        if m.group(1):
            flags = "bi" if m.group(1) == "***" else "b"
            out.extend((t, f + flags) for t, f in inline_spans(m.group(2)))
        elif m.group(3):
            out.extend((t, f + "i") for t, f in inline_spans(m.group(3)))
        else:
            out.append((m.group(4), "c"))
        pos = m.end()
    if pos < len(text):
        out.append((text[pos:], ""))
    return out


def add_text(paragraph, text: str, color: str) -> None:
    """Adds text with its **bold**, *italic* and `code` spans, and clickable links."""
    for part, flags in inline_spans(text):
        start = len(paragraph.runs)
        if "c" in flags:
            run = paragraph.add_run(part)
            _set_font(run._r, MONO)
            run.font.size = Pt(9.5)
            _shade_run(run, "F1F3F5")
        else:
            _add_linked(paragraph, part, color)
        for run in paragraph.runs[start:]:
            if "b" in flags:
                run.bold = True
            if "i" in flags:
                run.italic = True


MONO = "Liberation Mono"


def _shade_run(run, hex6: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex6)
    run._r.get_or_add_rPr().append(shd)


def _add_linked(paragraph, text: str, color: str) -> None:
    pos = 0
    for m in LINK.finditer(text):
        if m.start() > pos:
            paragraph.add_run(typo_fr(text[pos:m.start()]))
        target = m.group(0)
        url = target if target.startswith("http") else (f"mailto:{target}" if "@" in target else f"https://{target}")
        rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
        link = OxmlElement("w:hyperlink")
        link.set(qn("r:id"), rid)
        run = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), color)
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.extend([c, u])
        run.append(rpr)
        t = OxmlElement("w:t")
        t.text = target
        t.set(qn("xml:space"), "preserve")
        run.append(t)
        link.append(run)
        paragraph._p.append(link)
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(typo_fr(text[pos:]))


# --- document setup ----------------------------------------------------------

class Builder:
    def __init__(self, st: dict, images_dir: Path):
        self.st = st
        self.images_dir = images_dir
        self.doc = Document()
        self.first_in_section = True
        self.has_parts = False
        self.chapter_pages = True
        self._list_nums: dict[tuple[str, int], int] = {}
        self._last_list_key: tuple | None = None
        self._setup_page(self.doc.sections[0])
        self._setup_styles()

    # page & styles
    def _setup_page(self, section) -> None:
        top, bottom, left, right = self.st["margins"]
        section.page_width, section.page_height = Mm(210), Mm(297)
        section.top_margin, section.bottom_margin = Cm(top), Cm(bottom)
        section.left_margin, section.right_margin = Cm(left), Cm(right)
        section.footer_distance = Cm(1.2)

    @property
    def text_width(self) -> Emu:
        _, _, left, right = self.st["margins"]
        return Cm(21 - left - right)

    def _setup_styles(self) -> None:
        st, styles = self.st, self.doc.styles
        normal = styles["Normal"]
        _set_font(normal.element, st["font"])
        normal.font.size = Pt(st["size"])
        pf = normal.paragraph_format
        pf.line_spacing = st["line"]
        pf.space_after = Pt(6)
        pf.space_before = Pt(0)
        for i, size in enumerate(st["sizes"], start=1):
            h = styles[f"Heading {i}"]
            _set_font(h.element, st["font"])
            h.font.size = Pt(size)
            h.font.bold = True
            h.font.italic = i == 4
            _set_color(h.font, st["heading_color"])
            h.paragraph_format.space_before = Pt(18 if i == 1 else 14 if i == 2 else 10)
            h.paragraph_format.space_after = Pt(12 if i == 1 else 6)
            h.paragraph_format.line_spacing = 1.15
            h.paragraph_format.keep_with_next = True
            h.paragraph_format.first_line_indent = Cm(0)
            h.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        h1 = styles["Heading 1"]
        h1.font.all_caps = st["caps_h1"]
        if st["small_caps"]:
            h1.font.small_caps = styles["Heading 2"].font.small_caps = True
            h1.font.bold = False
        if st["center_h1"]:
            h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap = styles["Caption"]
        _set_font(cap.element, st["font"])
        cap.font.size = Pt(max(st["size"] - 1, 9))
        cap.font.bold = False
        cap.font.italic = False
        _set_color(cap.font, "000000")
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_before = Pt(6)
        cap.paragraph_format.space_after = Pt(6)
        cap.paragraph_format.first_line_indent = Cm(0)
        for name in ("List Bullet", "List Bullet 2", "List Bullet 3", "List Number", "List Number 2", "List Number 3"):
            ls = styles[name]
            _set_font(ls.element, st["font"])
            ls.paragraph_format.space_after = Pt(3)
            ls.paragraph_format.first_line_indent = None

    # structure
    def new_section(self, number_fmt: str | None, start: int | None) -> None:
        section = self.doc.add_section(WD_SECTION.NEW_PAGE)
        self._setup_page(section)
        if number_fmt:
            _page_numbers(section, number_fmt, start, self.st["page_number"], self.st["font"], "555555")
        self.first_in_section = True

    def page_break(self) -> None:
        if not self.first_in_section:
            self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        self.first_in_section = True

    def _para(self, text: str = "", style: str | None = None):
        p = self.doc.add_paragraph(style=style)
        if text:
            add_text(p, text, "1D4ED8" if self.st["theme"] != "moderne" else self.st["accent"])
        self.first_in_section = False
        if style is None or style == "Normal":
            pf = p.paragraph_format
            if self.st["justify"]:
                pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if self.st["indent"]:
                pf.first_line_indent = Cm(self.st["indent"])
        return p

    def liminal_title(self, text: str) -> None:
        """Titles of generated pages (Sommaire...) look like H1 but stay out of the TOC."""
        self.page_break()
        p = self._para()
        run = p.add_run(text.upper())
        run.bold = True
        run.font.size = Pt(self.st["sizes"][0])
        _set_color(run.font, self.st["heading_color"])
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.space_after = Pt(18)

    def marker(self, name: str) -> None:
        p = self._para(f"[[{name}]]")
        p.paragraph_format.first_line_indent = Cm(0)

    # blocks
    def heading(self, b: dict) -> None:
        level = max(1, min(int(b.get("level", 1)), 4))
        breaks = self.chapter_pages and (level == 1 or (level == 2 and self.has_parts and not b.get("special")))
        if breaks:
            self.page_break()
        if b.get("part"):
            # Part titles sit mid-page (space-before is ignored at a page top).
            for _ in range(10):
                self.doc.add_paragraph().paragraph_format.space_after = Pt(12)
        p = self.doc.add_paragraph(style=f"Heading {level}")
        p.add_run(b.get("text", ""))
        self.first_in_section = False
        if b.get("special") in NO_TOC_SPECIALS:
            p.style = self.doc.styles["Normal"]
            run = p.runs[0]
            run.bold = True
            run.font.size = Pt(self.st["sizes"][0])
            _set_color(run.font, self.st["heading_color"])
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if b.get("part"):
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # "Major" titles carry the theme's signature: chapters (level 2 under parts) or level 1.
        major = (level == 1 or (level == 2 and self.has_parts and not b.get("special"))) and not b.get("part")
        if major and b.get("special") not in NO_TOC_SPECIALS:
            if level == 2:  # promote chapter titles to the level-1 look
                for run in p.runs:
                    run.font.size = Pt(self.st["sizes"][0])
                    run.font.all_caps = self.st["caps_h1"]
                    run.font.small_caps = self.st["small_caps"]
                if self.st["center_h1"]:
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
            if self.st["rule_h1"]:
                _bottom_rule(p, self.st["accent"], 8 if self.st["small_caps"] else 12)
            if self.st["small_caps"]:
                _top_rule(p, self.st["accent"])
            if self.st["band_h1"]:
                _band(p, self.st["accent"])
                for run in p.runs:
                    _set_color(run.font, "FFFFFF")
        self._last_list_key = None

    def paragraph(self, b: dict) -> None:
        role = b.get("role")
        if role == "dedicace":
            p = self._para(b["text"])
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p.paragraph_format.first_line_indent = Cm(0)
            p.paragraph_format.left_indent = Cm(6)
            for r in p.runs:
                r.italic = True
        elif role == "biblio":
            p = self._para(b["text"])
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Cm(-1.25)
            p.paragraph_format.left_indent = Cm(1.25)
        elif role == "keywords":
            p = self._para()
            head, _, rest = b["text"].partition(":")
            p.add_run(head.strip() + " : ").bold = True
            p.add_run(rest.strip())
            p.paragraph_format.first_line_indent = Cm(0)
        else:
            self._para(b.get("text", ""))
        self._last_list_key = None

    def sigles(self, run: list[dict]) -> None:
        table = self.doc.add_table(rows=len(run), cols=2)
        _no_borders(table)
        table.autofit = False
        term_w = Cm(3.2)
        table.columns[0].width, table.columns[1].width = term_w, self.text_width - term_w
        for row, b in zip(table.rows, run):
            row.cells[0].width, row.cells[1].width = term_w, self.text_width - term_w
            t = row.cells[0].paragraphs[0].add_run(b.get("term", ""))
            t.bold = True
            row.cells[1].paragraphs[0].add_run(b.get("definition", ""))
            for c in row.cells:
                c.paragraphs[0].paragraph_format.space_after = Pt(2)
        self.first_in_section = False
        self._last_list_key = None

    def list_item(self, b: dict) -> None:
        level = max(0, min(int(b.get("level", 0)), 2))
        ordered = bool(b.get("ordered"))
        base = "List Number" if ordered else "List Bullet"
        style = base if level == 0 else f"{base} {level + 1}"
        p = self.doc.add_paragraph(style=style)
        add_text(p, b.get("text", ""), "1D4ED8")
        self.first_in_section = False
        if self.st["justify"]:
            p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if ordered:
            key = (style, level)
            if self._last_list_key is None or key not in self._list_nums:
                self._list_nums[key] = self._new_num(style)  # a new list starts at 1
            num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
            num_pr.get_or_add_numId().val = self._list_nums[key]
            num_pr.get_or_add_ilvl().val = 0
        self._last_list_key = "open"

    def _new_num(self, style_name: str) -> int:
        """A fresh numbering instance so each numbered list restarts at 1."""
        numbering = self.doc.part.numbering_part.element
        style_num_pr = self.doc.styles[style_name].element.pPr.numPr
        num_id = style_num_pr.numId.val
        abstract_id = None
        for num in numbering.findall(qn("w:num")):
            if num.get(qn("w:numId")) == str(num_id):
                abstract_id = num.find(qn("w:abstractNumId")).get(qn("w:val"))
        if abstract_id is None:
            return num_id
        new = numbering.add_num(int(abstract_id))
        new.add_lvlOverride(ilvl=0).add_startOverride(1)
        return new.numId

    def end_list(self) -> None:
        self._last_list_key = None

    def table(self, b: dict) -> None:
        rows = [list(r) for r in b.get("rows", []) if r]
        if not rows:
            return
        cols = max(len(r) for r in rows)
        rows = [r + [""] * (cols - len(r)) for r in rows]
        table = self.doc.add_table(rows=len(rows), cols=cols)
        table.style = self.doc.styles["Table Grid"]
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _full_width(table)
        for i, (row, values) in enumerate(zip(table.rows, rows)):
            for cell, value in zip(row.cells, values):
                p = cell.paragraphs[0]
                add_text(p, value, "1D4ED8")
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.first_line_indent = Cm(0)
                for run in p.runs:
                    run.font.size = Pt(max(self.st["size"] - 1, 9))
                    if i == 0:
                        run.bold = True
                        _set_color(run.font, self.st["table_head_text"])
                if i == 0:
                    _shade(cell, self.st["table_head"])
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Header row repeats on each page.
        trpr = table.rows[0]._tr.get_or_add_trPr()
        header = OxmlElement("w:tblHeader")
        header.set(qn("w:val"), "true")
        trpr.append(header)
        self.first_in_section = False
        self._para("").paragraph_format.space_after = Pt(0)
        self._last_list_key = None

    def figure(self, b: dict) -> None:
        path = self.images_dir / b.get("image", "")
        if not b.get("image") or not path.is_file():
            return
        with Image.open(path) as img:
            dpi = img.info.get("dpi", (96, 96))[0] or 96
            native = Emu(int(img.width / dpi * 914400))
        width = min(native, Emu(int(self.text_width * 0.9)))
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.keep_with_next = True
        p.add_run().add_picture(str(path), width=width)
        self.first_in_section = False
        self._last_list_key = None

    def caption(self, b: dict) -> None:
        label = "Tableau" if b.get("of") == "table" else "Figure"
        p = self.doc.add_paragraph(style="Caption")
        run = p.add_run(f"{label} ")
        run.bold = True
        _field(p, f"SEQ {label} \\* ARABIC")
        p.runs[-1].bold = True
        p.add_run("\u00a0: ").bold = True
        p.add_run(typo_fr(b.get("text", "")))
        if b.get("of") == "table":
            p.paragraph_format.keep_with_next = True
        self.first_in_section = False
        self._last_list_key = None

    def source(self, b: dict) -> None:
        p = self._para(b.get("text", ""))
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        for r in p.runs:
            r.italic = True
            r.font.size = Pt(max(self.st["size"] - 2, 8))
        self._last_list_key = None

    def quote(self, b: dict) -> None:
        p = self._para(b.get("text", ""))
        pf = p.paragraph_format
        pf.left_indent = pf.right_indent = Cm(1.5)
        pf.first_line_indent = Cm(0)
        pf.line_spacing = 1.0
        for r in p.runs:
            r.italic = True
        self._last_list_key = None

    def code(self, b: dict) -> None:
        """Code / command listing: monospace, grey box, lines kept as written, never split across pages if short."""
        lines = b.get("text", "").split("\n")
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            pf = p.paragraph_format
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
            pf.first_line_indent = Cm(0)
            pf.left_indent = Cm(0.4)
            pf.line_spacing = 1.0
            pf.space_before = Pt(6) if i == 0 else Pt(0)
            pf.space_after = Pt(8) if i == len(lines) - 1 else Pt(0)
            if len(lines) <= 25:
                pf.keep_with_next = i < len(lines) - 1
            run = p.add_run(line.replace("\t", "    ") or " ")
            _set_font(run._r, MONO)
            run.font.size = Pt(9)
            _set_color(run.font, "1F2937")
            _shade_para(p, "F3F4F6")
        self.first_in_section = False
        self._last_list_key = None

    def letterhead(self, h: dict) -> None:
        """Official Cameroonian header, as on exam papers and administrative letters:
        French column | logo | English column, groups separated by a short line of stars."""
        fr, en = h.get("fr") or [], h.get("en") or []
        if not fr and not en:
            return
        table = self.doc.add_table(rows=1, cols=3)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        _no_borders(table)
        _full_width(table)
        widths = [0.41, 0.18, 0.41]
        for i, w in enumerate(widths):
            table.columns[i].width = Emu(int(self.text_width * w))  # grid (LibreOffice)
            table.rows[0].cells[i].width = Emu(int(self.text_width * w))  # cell (Word)
        font_size = Pt(8.5)

        def column(cell, groups: list[list[str]]) -> None:
            first = True
            for gi, group in enumerate(groups):
                for line in group:
                    p = cell.paragraphs[0] if first else cell.add_paragraph()
                    first = False
                    pf = p.paragraph_format
                    pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    pf.first_line_indent = Cm(0)
                    pf.space_after = pf.space_before = Pt(0)
                    pf.line_spacing = 1.0
                    motto = bool(re.search(r"paix|peace", line, re.I))
                    run = p.add_run(line if motto else line.upper())
                    run.font.size = font_size
                    run.bold = not motto
                    run.italic = motto
                if gi < len(groups) - 1:
                    p = cell.add_paragraph()
                    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.space_after = Pt(0)
                    p.paragraph_format.first_line_indent = Cm(0)
                    run = p.add_run("*" * 8)
                    run.font.size = Pt(7)

        column(table.rows[0].cells[0], fr)
        column(table.rows[0].cells[2], en)
        mid = table.rows[0].cells[1]
        mid.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        p = mid.paragraphs[0]
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Cm(0)
        if h.get("logo_png"):
            p.add_run().add_picture(io.BytesIO(h["logo_png"]), height=Cm(2.3))
        spacer = self.doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(10)
        self.first_in_section = False
        self._last_list_key = None

    def title(self, b: dict) -> None:
        """Document title written at the top of the text (Markdown "# Title"), not part of the outline."""
        p = self._para()
        run = p.add_run(b.get("text", ""))
        pf = p.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
        if b.get("sub"):
            run.font.size = Pt(self.st["size"] + 1)
            run.italic = True
            _set_color(run.font, "555555")
            pf.space_after = Pt(18)
        else:
            run.bold = True
            run.font.size = Pt(self.st["sizes"][0] + 4)
            _set_color(run.font, self.st["heading_color"])
            pf.space_after = Pt(6)
            pf.keep_with_next = True
        self._last_list_key = None

    def cover(self, png: bytes) -> None:
        section = self.doc.sections[0]
        for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
            setattr(section, side, Mm(0))
        section.header_distance = section.footer_distance = Mm(0)
        p = self.doc.paragraphs[0] if self.doc.paragraphs else self.doc.add_paragraph()
        p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(0)
        p.add_run().add_picture(io.BytesIO(png), width=Mm(209.5))
        self.first_in_section = False


# --- orchestration -------------------------------------------------------------

def build_docx(document: dict, images_dir: Path) -> bytes:
    st = resolved_style(document.get("style"))
    opts = {**DEFAULT_OPTIONS, **(document.get("options") or {})}
    if "chapter_pages" not in (document.get("options") or {}):  # documents saved before this option existed
        opts["chapter_pages"] = KINDS.get((document.get("meta") or {}).get("kind", ""), {}).get("chapter_pages", False)
    blocks = [b for b in document.get("blocks", []) if not b.get("hidden")]
    b = Builder(st, images_dir)
    b.has_parts = any(x.get("part") for x in blocks)
    b.chapter_pages = opts["chapter_pages"]

    started = False
    has_cover = opts["cover"] and document.get("cover_svg")
    # The document's own title (and subtitle) above the first heading always opens the document:
    # on the cover when there is one (so it is not repeated), otherwise at the top of page 1, before any sommaire.
    first_heading = next((i for i, x in enumerate(blocks) if x.get("type") == "heading"), len(blocks))
    titles = [x for x in blocks[:first_heading] if x.get("type") == "title"]
    blocks = [x for x in blocks if not (x.get("type") == "title" and any(x is t for t in titles))]
    if opts.get("letterhead") and document.get("letterhead") and not has_cover:
        b.letterhead(_letterhead_data(document["letterhead"]))
    if has_cover:
        b.cover(cover_render.svg_to_png(document["cover_svg"], 200))
        started = True
    elif titles:
        for t in titles:
            b.title(t)
        b.first_in_section = True  # the sommaire or the text follows on the same page

    # Split: preliminary blocks (before the first non-preliminary heading) / body.
    split = next(
        (i for i, x in enumerate(blocks)
         if x.get("type") == "heading" and x.get("special") not in PRELIM_SPECIALS),
        len(blocks),
    )
    prelim, body = blocks[:split], blocks[split:]
    has_tables = any(x.get("type") == "caption" and x.get("of") == "table" for x in blocks)
    has_figures = any(x.get("type") == "caption" and x.get("of") == "figure" for x in blocks)
    # Front matter (roman page numbers) only when there is some: dedication, thanks, abstract... or a sommaire.
    # An introduction paragraph above the first heading of a course is just the start of the text.
    want_prelim = any(x.get("special") for x in prelim) or opts["toc"] or (opts["lists"] and (has_tables or has_figures))
    if not want_prelim:
        prelim, body = [], prelim + body
    number = opts["page_numbers"]

    if want_prelim:
        if started:
            b.new_section("lowerRoman" if number else None, 1)
        elif number:
            _page_numbers(b.doc.sections[0], "lowerRoman", 1, st["page_number"], st["font"], "555555")
        started = True
        if opts["toc"]:
            b.liminal_title("Sommaire")
            b.marker("TOC:2")
        lists_done = False
        for x in _group_sigles(prelim):
            if not lists_done and x.get("type") == "heading" and x.get("special") in ("resume", "abstract"):
                _lists(b, opts, has_tables, has_figures)
                lists_done = True
            _emit(b, x, prelim)
        if not lists_done:
            _lists(b, opts, has_tables, has_figures)

    if started:
        b.new_section("decimal" if number else None, 1)
    elif number:
        _page_numbers(b.doc.sections[0], "decimal", 1, st["page_number"], st["font"], "555555")
    for x in _group_sigles(body):
        _emit(b, x, body)

    if opts["toc_end"]:
        b.liminal_title("Table des matières")
        b.marker("TOC:4")

    buf = io.BytesIO()
    b.doc.save(buf)
    return buf.getvalue()


def _letterhead_data(h: dict) -> dict:
    """Decodes the embedded logo (data: URL) into image bytes Word can hold."""
    out = {"fr": h.get("fr") or [], "en": h.get("en") or []}
    logo = h.get("logo") or ""
    m = re.match(r"^data:image/(png|jpeg|jpg|webp);base64,(.+)$", logo, re.S)
    if m:
        try:
            raw = base64.b64decode(m.group(2))
            with Image.open(io.BytesIO(raw)) as img:
                buf = io.BytesIO()
                img.convert("RGBA").save(buf, "PNG")  # webp/odd PNGs → plain PNG
                out["logo_png"] = buf.getvalue()
        except Exception:  # noqa: BLE001 - a broken logo must not break the document
            pass
    return out


def _lists(b: Builder, opts: dict, has_tables: bool, has_figures: bool) -> None:
    if not opts["lists"]:
        return
    if has_tables:
        b.liminal_title("Liste des tableaux")
        b.marker("LOT")
    if has_figures:
        b.liminal_title("Liste des figures")
        b.marker("LOF")


def _group_sigles(blocks: list[dict]) -> list[dict]:
    out, run = [], []
    for x in blocks:
        if x.get("role") == "sigle":
            run.append(x)
            continue
        if run:
            out.append({"type": "_sigles", "items": run})
            run = []
        out.append(x)
    if run:
        out.append({"type": "_sigles", "items": run})
    return out


def _emit(b: Builder, x: dict, _all: list[dict]) -> None:
    kind = x.get("type")
    if kind != "list":
        b.end_list()
    if kind == "heading":
        b.heading(x)
    elif kind == "paragraph":
        if x.get("role") == "sigle":
            b.sigles([x])
        else:
            b.paragraph(x)
    elif kind == "_sigles":
        b.sigles(x["items"])
    elif kind == "list":
        b.list_item(x)
    elif kind == "table":
        b.table(x)
    elif kind == "figure":
        b.figure(x)
    elif kind == "caption":
        b.caption(x)
    elif kind == "source":
        b.source(x)
    elif kind == "quote":
        b.quote(x)
    elif kind == "code":
        b.code(x)
    elif kind == "title":
        b.title(x)
