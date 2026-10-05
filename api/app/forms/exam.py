"""Teacher's exam paper (épreuve), bilingual MINESEC-style letterhead (French | English).

The teacher types (or pastes) the exercises as plain text; we recognise:
- "Exercice 1 (5 pts)", "Partie A : …", "Problème (8 points)" -> section titles, points right-aligned
- "1) …", "2. …" -> numbered questions (hanging indent); "(2 pts)" at the end -> points right-aligned
- "a) …" -> sub-questions; "A) …" -> multiple-choice options
"""
from __future__ import annotations

import io
import re

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .common import field, new_document, no_borders, para

SECTION = re.compile(
    r"^(exercice|exercise|partie|part|probl[èe]me|problem|section|situation\s+probl[èe]me)\b\s*([ivx\d]+|[a-e])?\s*[:.\-–]?\s*(.*)$",
    re.I,
)
POINTS = re.compile(r"\(?\s*(\d+(?:[.,]\d+)?)\s*(pts?|points?)\s*\)?\s*\.?$", re.I)
QUESTION = re.compile(r"^(\d{1,2})\s*[).\-]\s+(.*)$")
SUBQ = re.compile(r"^([a-h])\s*[).]\s+(.*)$")
OPTION = re.compile(r"^([A-E])\s*[).]\s+(.*)$")

TEXT_W = Cm(21 - 2 - 2)


def _s(data: dict, key: str, default: str = "") -> str:
    return str(data.get(key) or default).strip()[:400]


def split_points(text: str) -> tuple[str, str]:
    m = POINTS.search(text)
    if not m or m.start() == 0:
        return text, ""
    pts = m.group(1).replace(",", ".")
    return text[: m.start()].rstrip(" -–:,"), f"{pts.rstrip('0').rstrip('.') if '.' in pts else pts} pt{'s' if float(pts) > 1 else ''}"


def parse(content: str) -> list[dict]:
    items = []
    for raw in content.replace("\r\n", "\n").split("\n"):
        line = raw.strip()
        if not line:
            continue
        m = SECTION.match(line)
        if m and len(line) < 140:
            body, pts = split_points(line)
            items.append({"kind": "section", "text": body, "points": pts})
            continue
        for kind, rx in (("question", QUESTION), ("sub", SUBQ), ("option", OPTION)):
            m = rx.match(line)
            if m:
                body, pts = split_points(m.group(2))
                items.append({"kind": kind, "label": m.group(1), "text": body, "points": pts})
                break
        else:
            body, pts = split_points(line)
            items.append({"kind": "text", "text": body, "points": pts})
    return items


def total_points(items: list[dict]) -> float:
    sections = [i for i in items if i["kind"] == "section" and i["points"]]
    source = sections or [i for i in items if i["points"]]
    total = 0.0
    for i in source:
        try:
            total += float(i["points"].split()[0])
        except ValueError:
            pass
    return total


def _with_points(p, text: str, points: str, bold=False) -> None:
    p.paragraph_format.tab_stops.add_tab_stop(TEXT_W, WD_TAB_ALIGNMENT.RIGHT)
    run = p.add_run(text)
    run.bold = bold
    if points:
        pr = p.add_run(f"\t({points})")
        pr.bold = True
        pr.font.size = Pt(10.5)


def build_exam(data: dict) -> bytes:
    doc = new_document("Times New Roman", 12, margins=(1.6, 1.6, 2, 2))
    color = _s(data, "color", "0E9F6E").lstrip("#").upper()

    items = parse(str(data.get("content") or "")[:60000])

    # Letterhead, as on the papers handed out in Cameroonian schools:
    # French (left) | English (right), then the year on the right, then the exam facts on the left.
    bilingual = data.get("bilingual", True) is not False
    ministry = _s(data, "ministry", "MINISTÈRE DES ENSEIGNEMENTS SECONDAIRES")
    ministry_en = _s(data, "ministry_en") or ("MINISTRY OF SECONDARY EDUCATION" if "SECONDAIRE" in ministry.upper() else "")
    school = _s(data, "school")
    fr = []
    if data.get("country", True):
        fr += [("RÉPUBLIQUE DU CAMEROUN", True, False), ("Paix – Travail – Patrie", False, True)]
    fr += [(ministry, True, False)]
    for key in ("delegation", "department"):
        if _s(data, key):
            fr.append((_s(data, key), False, False))
    if school:
        fr.append((school, True, False))
    en = []
    if data.get("country", True):
        en += [("REPUBLIC OF CAMEROON", True, False), ("Peace – Work – Fatherland", False, True)]
    if ministry_en:
        en.append((ministry_en, True, False))
    if _s(data, "delegation_en"):
        en.append((_s(data, "delegation_en"), False, False))
    if school:
        en.append((_s(data, "school_en") or school, True, False))

    head = doc.add_table(rows=1, cols=2 if bilingual else 1)
    no_borders(head)
    head.autofit = False
    cells = head.rows[0].cells
    for cell, lines in zip(cells, (fr, en)):
        cell.width = Cm(8.5 if bilingual else 17)
        first = True
        for text, bold, italic in lines:
            p = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            run = p.add_run(text)
            run.bold, run.italic = bold, italic
            run.font.size = Pt(9.5)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if _s(data, "year"):
        p = para(doc, "", space_before=10, align="right")
        p.add_run("Année scolaire : ").font.size = Pt(11)
        p.add_run(_s(data, "year")).font.size = Pt(11)

    facts = [("Classe", _s(data, "class")), ("Durée", _s(data, "duration")),
             ("Coefficient", _s(data, "coef")), ("Examinateur", _s(data, "teacher"))]
    first = True
    for label, value in facts:
        if value:
            p = para(doc, "", space_before=10 if first else 0)
            first = False
            r = p.add_run(f"{label} : {value}")
            r.bold = True
            r.font.size = Pt(11)

    # Title: "ÉPREUVE DE COUPE, 20 pts" (the total is the one written on the paper, else the sum of the exercises)
    if _s(data, "exam"):
        para(doc, _s(data, "exam"), size=10.5, align="center", space_before=10)
    total = _s(data, "total").replace(",", ".")
    try:
        points = float(total) if total else total_points(items)
    except ValueError:
        points = total_points(items)
    title = f"ÉPREUVE DE {_s(data, 'subject', 'MATIÈRE').upper()}"
    if points:
        shown = f"{points:g}"
        title += f", {shown} pt{'s' if points > 1 else ''}"
    p = para(doc, "", align="center", space_before=10, space_after=6)
    r = p.add_run(title)
    r.bold = r.underline = True
    r.font.size = Pt(14)

    if _s(data, "instructions"):
        para(doc, _s(data, "instructions"), italic=True, size=10.5, align="center", space_before=8)

    # Body
    for it in items:
        k = it["kind"]
        if k == "section":
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            _with_points(p, it["text"].upper(), it["points"], bold=True)
            for run in p.runs[:1]:
                run.font.color.rgb = RGBColor.from_string(color)
            bdr = OxmlElement("w:pBdr")
            bottom = OxmlElement("w:bottom")
            for key, val in (("w:val", "single"), ("w:sz", "6"), ("w:space", "2"), ("w:color", color)):
                bottom.set(qn(key), val)
            bdr.append(bottom)
            p._p.get_or_add_pPr().append(bdr)
        elif k in ("question", "sub", "option"):
            p = doc.add_paragraph()
            indent = {"question": 0.8, "sub": 1.6, "option": 1.6}[k]
            p.paragraph_format.left_indent = Cm(indent)
            p.paragraph_format.first_line_indent = Cm(-0.7)
            p.paragraph_format.space_after = Pt(4)
            label = {"question": f"{it['label']}.", "sub": f"{it['label']})", "option": f"{it['label']}."}[k]
            lab = p.add_run(f"{label} ")
            lab.bold = k == "question"
            _with_points(p, it["text"], it["points"])
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if len(it["text"]) > 60:  # a paragraph of text (the statement), not a short line
                p.paragraph_format.first_line_indent = Cm(1)
            _with_points(p, it["text"], it["points"])

    # Footer: page x / y
    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run(f"{_s(data, 'subject', 'Épreuve')} · {_s(data, 'class')} · Page ")
    run.font.size = Pt(9)
    field(footer, "PAGE")
    footer.add_run(" / ").font.size = Pt(9)
    field(footer, "NUMPAGES")
    for r in footer.runs:
        r.font.size = Pt(9)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
