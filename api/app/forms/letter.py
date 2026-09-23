"""Administrative letter / request, Cameroonian layout.

    NOM Prénom                              Yaoundé, le 23 septembre 2026
    Adresse · Tél · E-mail                                    [Timbre fiscal]
                                            À
                                            Monsieur le Directeur Général
                                            de ENEO Cameroun, Douala
                                            s/c de Monsieur le Chef du personnel
    Objet : Demande d'emploi
    P.J. : CV ; copie du diplôme

            Monsieur le Directeur Général,
            J'ai l'honneur de venir très respectueusement…
                                                          Signature
                                                          NOM Prénom
"""
from __future__ import annotations

import io

from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Cm, Pt

from .common import boxed, new_document, no_borders, para

MAX = 4000


def _s(data: dict, key: str, default: str = "") -> str:
    return str(data.get(key) or default).strip()[:MAX]


def salutation(recipient: dict) -> str:
    civ = _s(recipient, "civility", "Monsieur")
    title = _s(recipient, "title")
    return f"{civ} {title}".strip() + ","


def build_letter(data: dict) -> bytes:
    doc = new_document("Times New Roman", 12)
    sender = data.get("sender") or {}
    recipient = data.get("recipient") or {}

    # Sender (left) | place, date and optional tax stamp (right)
    top = doc.add_table(rows=1, cols=2)
    no_borders(top)
    top.autofit = False
    left, right = top.rows[0].cells
    left.width, right.width = Cm(9.5), Cm(6.8)
    left.paragraphs[0].add_run(_s(sender, "name", "NOM Prénom")).bold = True
    for key in ("extra", "address", "phone", "email"):
        value = _s(sender, key)
        if value:
            label = {"phone": "Tél. : ", "email": "E-mail : "}.get(key, "")
            para(left, f"{label}{value}", size=11)
    place_date = ", le ".join(x for x in (_s(data, "place"), _s(data, "date")) if x)
    rp = right.paragraphs[0]
    rp.add_run(place_date)
    rp.alignment = 2  # right
    if data.get("stamp"):
        stamp = right.add_table(rows=1, cols=1)
        stamp.alignment = WD_TABLE_ALIGNMENT.RIGHT
        cell = stamp.rows[0].cells[0]
        cell.width = Cm(3)
        boxed(cell)
        p = cell.paragraphs[0]
        p.alignment = 1
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(14)
        run = p.add_run("Timbre fiscal")
        run.italic = True
        run.font.size = Pt(9)

    # Recipient block, pushed to the right
    para(doc, "", space_after=18)
    lines = ["À", f"{_s(recipient, 'civility', 'Monsieur')} {_s(recipient, 'title')}".strip()]
    if _s(recipient, "org"):
        lines.append(_s(recipient, "org"))
    if _s(recipient, "city"):
        lines.append(_s(recipient, "city"))
    for i, line in enumerate(lines):
        p = para(doc, line, bold=i == 1)
        p.paragraph_format.left_indent = Cm(9)
    if _s(data, "via"):
        p = para(doc, f"s/c de {_s(data, 'via')}", italic=True, space_before=6)
        p.paragraph_format.left_indent = Cm(9)

    # Subject and attachments
    para(doc, "", space_after=14)
    p = doc.add_paragraph()
    p.add_run("Objet : ").bold = True
    subject = p.add_run(_s(data, "subject", "Demande"))
    subject.bold = True
    subject.underline = True
    attachments = [str(a).strip() for a in (data.get("attachments") or []) if str(a).strip()]
    if attachments:
        p = doc.add_paragraph()
        p.add_run("P.J. : ").bold = True
        p.add_run(" ; ".join(attachments[:12]))
        p.paragraph_format.space_before = Pt(4)

    # Salutation + body
    p = para(doc, salutation(recipient), space_before=20, space_after=10)
    p.paragraph_format.first_line_indent = Cm(1.25)
    body = _s(data, "body").replace("\r\n", "\n")
    for chunk in [c.strip() for c in body.split("\n\n") if c.strip()]:
        p = para(doc, " ".join(chunk.split("\n")), align="justify", space_after=8)
        p.paragraph_format.first_line_indent = Cm(1.25)
        p.paragraph_format.line_spacing = 1.3

    # Signature
    para(doc, "", space_after=10)
    sig = para(doc, _s(data, "signature_label", "Signature"), italic=True, size=10, align="right")
    sig.paragraph_format.space_after = Pt(36)
    para(doc, _s(sender, "name", "NOM Prénom"), bold=True, align="right")

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
