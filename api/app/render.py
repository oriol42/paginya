"""Cover rendering: sanitized SVG -> PNG (300 dpi) -> PDF and DOCX.

For the MVP the cover is embedded in the Word file as a full-page image:
it prints exactly like the preview. Editable text boxes can come later.
"""
import io
import re

import img2pdf
import resvg_py
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.shared import Mm
from PIL import Image

from .config import FONTS_DIR

A4_MM = (210, 297)
A4_PT_WIDTH = 595.28  # covers are drawn in an A4 viewBox of 595 x 842 points
PRINT_DPI = 300
PREVIEW_DPI = 110

WATERMARK = """
<g opacity="0.16" font-family="Poppins" font-weight="800" font-size="46" fill="#065F46"
   text-anchor="middle" transform="rotate(-35 297.5 421)">
  <text x="297.5" y="260">APERÇU · PAGINYA</text>
  <text x="297.5" y="440">APERÇU · PAGINYA</text>
  <text x="297.5" y="620">APERÇU · PAGINYA</text>
</g>
"""


def add_watermark(svg: str) -> str:
    return re.sub(r"</svg>\s*$", WATERMARK + "</svg>", svg.strip())


def svg_to_png(svg: str, dpi: int = PRINT_DPI) -> bytes:
    zoom = dpi / 72  # the SVG user unit is one typographic point
    png = resvg_py.svg_to_bytes(
        svg_string=svg,
        font_dirs=[str(FONTS_DIR)],
        skip_system_fonts=True,
        zoom=zoom,
        background="#ffffff",
    )
    return bytes(png)


def png_to_pdf(png: bytes) -> bytes:
    # img2pdf refuses images with an alpha channel: flatten to RGB first.
    img = Image.open(io.BytesIO(png)).convert("RGB")
    buf = io.BytesIO()
    img.save(buf, format="JPEG", quality=92, dpi=(PRINT_DPI, PRINT_DPI))
    layout = img2pdf.get_layout_fun((img2pdf.mm_to_pt(A4_MM[0]), img2pdf.mm_to_pt(A4_MM[1])))
    return img2pdf.convert(buf.getvalue(), layout_fun=layout)


def png_to_docx(png: bytes) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(A4_MM[0]), Mm(A4_MM[1])
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, side, Mm(0))
    section.header_distance = section.footer_distance = Mm(0)
    par = doc.paragraphs[0] if doc.paragraphs else doc.add_paragraph()
    par.paragraph_format.space_before = par.paragraph_format.space_after = 0
    # Slightly smaller than the page so Word never pushes it onto a 2nd page.
    par.add_run().add_picture(io.BytesIO(png), width=Mm(A4_MM[0] - 0.5))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def render_cover(svg: str, fmt: str, watermark: bool) -> bytes:
    if watermark:
        svg = add_watermark(svg)
    if fmt == "preview":
        return svg_to_png(svg, PREVIEW_DPI)
    png = svg_to_png(svg, PRINT_DPI)
    if fmt == "png":
        return png
    if fmt == "pdf":
        return png_to_pdf(png)
    if fmt == "docx":
        return png_to_docx(png)
    raise ValueError(f"Format inconnu : {fmt}")
