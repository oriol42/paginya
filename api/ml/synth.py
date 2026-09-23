"""Synthetic training data for structure detection: well-made documents → realistically messy copies.

1. `truth_doc()` builds the logical structure of a Cameroonian academic/professional document
   (specials, parts, chapters, sections, paragraphs, lists, captions, typed table of contents).
2. A random "author habit" renders it the way students really do: Word heading styles or manual
   bold/size/caps, typed numbering (CHAPITRE I, I., A., 1., 1.1...), typed bullets, emphasis in bold,
   inconsistent sizes, missing formatting... as a DOCX, or as pasted plain text.
3. The file goes through the *real* Paginya extraction, so the model learns from exactly what the
   app sees in production. Each extracted paragraph keeps the label of the element it came from.
"""
from __future__ import annotations

import random
import tempfile
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.doc.extract import Raw, from_docx, from_text

# Labels
PARA, H1, H2, H3, H4, LIST, CAPTION, TOC = range(8)
NAMES = ["paragraphe", "titre 1", "titre 2", "titre 3", "titre 4", "liste", "légende", "sommaire"]


@dataclass
class El:
    label: int
    text: str
    depth: int = 0        # heading depth in the tree (1..4)
    kind: str = ""        # special | partie | chapitre | section
    ordered: bool = False  # list


# ----------------------------------------------------------------- vocabulary

ORGS = ["ADVANS Cameroun", "la CAMTEL", "ENEO", "la mairie de Douala 5e", "MTN Cameroon", "l'hôpital Laquintinie", "la SOSUCAM",
        "Afriland First Bank", "le MINEPAT", "la CNPS", "Orange Cameroun", "le Port Autonome de Douala", "l'ESSTIC", "la CRTV",
        "Express Union", "la Délégation régionale du Commerce", "une PME de transformation de manioc", "la BICEC"]
TOPICS = ["la gestion des stocks", "la communication interne", "le recouvrement des crédits", "la satisfaction des clients",
          "la digitalisation des services", "la gestion des ressources humaines", "la comptabilité analytique", "le marketing digital",
          "la maintenance du réseau", "la sécurité informatique", "la chaîne d'approvisionnement", "l'accueil des usagers",
          "l'archivage des documents", "la politique tarifaire", "la production agricole", "le suivi des projets"]
SECTION_TITLES = [
    "Historique", "Missions et objectifs", "Organisation et fonctionnement", "Situation géographique", "Organigramme",
    "Les activités de la structure", "Déroulement du stage", "Les tâches effectuées", "Les difficultés rencontrées",
    "Analyse de l'existant", "Méthodologie de recherche", "Présentation des résultats", "Discussion des résultats",
    "Les recommandations", "Le cadre théorique", "Le cadre conceptuel", "La revue de la littérature", "Les limites de l'étude",
    "Les forces et faiblesses", "Les acquis du stage", "Présentation du service d'accueil", "Les outils utilisés",
    "La collecte des données", "Le traitement des données", "Les perspectives", "Les contraintes budgétaires",
    "Le personnel", "Les partenaires", "Les produits et services", "La clientèle cible", "La concurrence",
]
CHAPTER_TITLES = ["Présentation de la structure d'accueil", "Cadre théorique et conceptuel", "Déroulement du stage",
                  "Analyse et diagnostic", "Méthodologie de l'étude", "Présentation et analyse des résultats",
                  "Suggestions et recommandations", "Historique et missions", "Organisation administrative"]
PART_TITLES = ["Présentation de la structure d'accueil", "Déroulement du stage", "Cadre théorique et méthodologique",
               "Analyse des résultats et recommandations", "Le cadre de l'étude", "Le travail effectué"]
SUBJ = ["Le service", "Notre stage", "Cette structure", "L'entreprise", "Le personnel", "La direction", "Ce travail", "L'étude",
        "Le responsable du service", "Chaque agent", "Le département", "La cellule informatique", "Notre encadreur"]
VERB = ["a permis de", "vise à", "contribue à", "est chargé de", "s'occupe de", "doit assurer", "cherche à", "a pour rôle de",
        "a été créé pour", "participe à", "nous a appris à", "consiste à"]
OBJ = ["améliorer la qualité du service", "renforcer la relation avec la clientèle", "suivre les opérations quotidiennes",
       "réduire les délais de traitement", "former les nouveaux agents", "analyser les données collectées", "moderniser les outils",
       "garantir la sécurité des opérations", "répondre aux besoins des usagers", "optimiser les ressources disponibles",
       "mettre en place de nouvelles procédures", "produire des rapports mensuels"]
TAIL = ["dans un contexte de forte concurrence", "conformément aux textes en vigueur", "depuis sa création en {y}",
        "avec l'appui de ses partenaires", "au niveau national", "dans la ville de {c}", "sous la supervision du chef de service",
        "grâce à une équipe dynamique", "malgré des moyens limités", "pour la période allant de {m1} à {m2}"]
CITIES = ["Douala", "Yaoundé", "Bafoussam", "Garoua", "Bamenda", "Ngaoundéré", "Kribi", "Limbé", "Bertoua", "Ebolowa"]
MONTHS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre", "octobre", "novembre", "décembre"]
LIST_ITEMS = ["l'accueil et l'orientation des clients", "la saisie des opérations", "le classement des dossiers", "la rédaction des courriers",
              "la participation aux réunions", "le suivi des indicateurs", "la mise à jour de la base de données", "l'élaboration des rapports",
              "la vérification des pièces justificatives", "l'assistance aux usagers", "la gestion des réclamations", "l'inventaire du matériel",
              "la formation des stagiaires", "la préparation des budgets", "la communication sur les réseaux sociaux"]
EMPH = ["NB : les données présentées sont celles de l'année {y}.", "Il convient de noter que ce service est le plus sollicité.",
        "Source : enquête de terrain, {y}.", "Pour ce faire :", "Ainsi :", "À cet effet, plusieurs actions ont été menées :"]

ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X"]
ORD_FR = ["PREMIÈRE", "DEUXIÈME", "TROISIÈME", "QUATRIÈME"]


# Out-of-distribution check: SYNTH_SPLIT=train uses half of the vocabulary, =test the other half
# (and real sentences from the fixture for the body), so we measure generalisation, not memory.
import os  # noqa: E402

_SPLIT = os.environ.get("SYNTH_SPLIT", "")
if _SPLIT:
    keep = 0 if _SPLIT == "train" else 1
    SECTION_TITLES = SECTION_TITLES[keep::2]
    CHAPTER_TITLES = CHAPTER_TITLES[keep::2]
    PART_TITLES = PART_TITLES[keep::2]
    LIST_ITEMS = LIST_ITEMS[keep::2]
    OBJ = OBJ[keep::2]
_REAL = []
if _SPLIT == "test":
    import re as _re

    _fx = Path(__file__).resolve().parent.parent / "tests" / "fixtures" / "rapport_brut.txt"
    _REAL = [s.strip() + "." for s in _re.split(r"(?<=[.!?])\s+", _fx.read_text()) if len(s) > 60 and not s.isupper()]


def sentence(rng: random.Random) -> str:
    if _REAL and rng.random() < 0.7:
        return rng.choice(_REAL)
    t = rng.choice(TAIL).format(y=rng.randint(1995, 2025), c=rng.choice(CITIES), m1=rng.choice(MONTHS), m2=rng.choice(MONTHS))
    s = f"{rng.choice(SUBJ)} {rng.choice(VERB)} {rng.choice(OBJ)} {t}"
    if rng.random() < 0.3:
        s += f", notamment pour {rng.choice(TOPICS)}"
    return s + "."


def paragraph(rng: random.Random) -> str:
    n = rng.choices([1, 2, 3, 4, 5, 6], [8, 18, 26, 24, 14, 10])[0]
    return " ".join(sentence(rng) for _ in range(n))


# ----------------------------------------------------------------- truth

def truth_doc(rng: random.Random) -> list[El]:
    """Logical structure (what a well-made document contains)."""
    els: list[El] = []
    with_parts = rng.random() < 0.45
    add_toc = rng.random() < 0.55
    parts = rng.randint(2, 3) if with_parts else 1
    chapters_per_part = rng.randint(1, 3)

    def special(name: str) -> None:
        els.append(El(H1, name, 1, "special"))

    if rng.random() < 0.5:
        special(rng.choice(["DÉDICACE", "Dédicace", "DEDICACES"]))
        els.append(El(PARA, f"À mes parents {rng.choice(['MBARGA', 'NJOYA', 'TCHOUA', 'ATANGANA', 'FOTSO'])} et à toute ma famille."))
    if rng.random() < 0.7:
        special(rng.choice(["REMERCIEMENTS", "Remerciements"]))
        els.append(El(PARA, paragraph(rng)))
        for _ in range(rng.randint(2, 5)):
            els.append(El(LIST, f"{rng.choice(['Au', 'À'])} {rng.choice(['Directeur', 'chef de service', 'Professeur', 'personnel'])} de {rng.choice(ORGS)} pour son accueil ;"))
    if add_toc:
        els.append(El(H1, rng.choice(["SOMMAIRE", "TABLE DES MATIÈRES", "Sommaire"]), 1, "special"))
        els[-1].label = TOC  # the TOC title itself is dropped by Paginya (it rebuilds the table)
        for t in ["Introduction", "Chapitre I", "Chapitre II", "Conclusion", "Bibliographie"][: rng.randint(3, 5)]:
            els.append(El(TOC, f"{t} {'.' * rng.randint(8, 40)} {rng.randint(1, 60)}"))
    special(rng.choice(["INTRODUCTION GÉNÉRALE", "INTRODUCTION", "Introduction générale"]))
    for _ in range(rng.randint(1, 4)):
        els.append(El(PARA, paragraph(rng)))

    ch = 0
    for p in range(parts):
        if with_parts:
            els.append(El(H1, rng.choice(PART_TITLES), 1, "partie"))
            els[-1].text = f"{ORD_FR[p]} PARTIE : {els[-1].text}"
        for _ in range(chapters_per_part):
            ch += 1
            d = 2 if with_parts else 1
            els.append(El([H1, H2][d - 1], rng.choice(CHAPTER_TITLES), d, "chapitre"))
            els[-1].text = f"CHAPITRE {ROMAN[ch - 1]} : {els[-1].text}"
            if rng.random() < 0.5:
                els.append(El(PARA, paragraph(rng)))
            for _s in range(rng.randint(2, 4)):
                d2 = d + 1
                els.append(El([H1, H2, H3, H4][d2 - 1], rng.choice(SECTION_TITLES), d2, "section"))
                for _k in range(rng.randint(1, 3)):
                    els.append(El(PARA, paragraph(rng)))
                if rng.random() < 0.45 and d2 < 4:
                    for _ss in range(rng.randint(2, 3)):
                        d3 = d2 + 1
                        els.append(El([H1, H2, H3, H4][d3 - 1], rng.choice(SECTION_TITLES), d3, "section"))
                        els.append(El(PARA, paragraph(rng)))
                        if rng.random() < 0.3 and d3 < 4:
                            els.append(El(H4, rng.choice(SECTION_TITLES), 4, "section"))
                            els.append(El(PARA, paragraph(rng)))
                if rng.random() < 0.2:
                    # enumeration typed WITHOUT bullets: "… les suivantes :" then short lines
                    els.append(El(PARA, rng.choice(["Les missions principales sont les suivantes :", "Nos tâches étaient les suivantes :",
                                                   "Ce service comprend :", "Les objectifs sont :", "Il s'agit notamment de :"])))
                    for _ in range(rng.randint(2, 5)):
                        it = rng.choice(LIST_ITEMS)
                        els.append(El(LIST, (it[:1].upper() + it[1:]) + rng.choice(["", "", " ;", "."]), ordered=False, kind="nomark"))
                if rng.random() < 0.35:
                    if rng.random() < 0.5:
                        els.append(El(PARA, rng.choice(EMPH).format(y=rng.randint(2015, 2025))))
                    ordered = rng.random() < 0.3
                    for _ in range(rng.randint(2, 6)):
                        els.append(El(LIST, rng.choice(LIST_ITEMS) + rng.choice([" ;", ";", ".", ""]), ordered=ordered))
                if rng.random() < 0.25:
                    kind = rng.choice(["Tableau", "Figure", "Graphique", "Tab.", "Fig."])
                    els.append(El(CAPTION, f"{kind} {rng.randint(1, 12)} : {rng.choice(SECTION_TITLES)} de {rng.choice(ORGS)}"))
                    if rng.random() < 0.5:
                        els.append(El(PARA, f"Source : {rng.choice(['nos enquêtes', 'service comptable', 'rapport annuel'])} {rng.randint(2018, 2025)}"))

    special(rng.choice(["CONCLUSION GÉNÉRALE", "CONCLUSION", "Conclusion générale", "CONCLUSION ET RECOMMANDATIONS"]))
    for _ in range(rng.randint(1, 3)):
        els.append(El(PARA, paragraph(rng)))
    special(rng.choice(["BIBLIOGRAPHIE", "RÉFÉRENCES BIBLIOGRAPHIQUES", "Bibliographie", "WEBOGRAPHIE"]))
    for _ in range(rng.randint(2, 5)):
        els.append(El(PARA, f"{rng.choice(['NGONO', 'FOUDA', 'KAMGA', 'ESSOMBA', 'BELLO'])} {rng.choice('ABCDJMP')}. ({rng.randint(1998, 2024)}), {rng.choice(CHAPTER_TITLES)}, {rng.choice(CITIES)}, Presses universitaires."))
    if rng.random() < 0.4:
        special(rng.choice(["ANNEXES", "Annexes"]))
    return els


# ----------------------------------------------------------------- author habits

NUMBERINGS = {
    # heading depth → how the student numbers it (after chapters / parts, which carry their own words)
    "roman_letter": ["", "{R}.", "{A}.", "{n}."],
    "roman_letter2": ["", "{R}-", "{A}-", "{n})"],
    "decimal": ["", "{n}.", "{n1}.{n}.", "{n1}.{n2}.{n}."],
    "decimal_nodot": ["", "{n}", "{n1}.{n}", "{n1}.{n2}.{n}"],
    "none": ["", "", "", ""],
    "mixed": ["", "{R}.", "{n}.", "{a})"],
}


@dataclass
class Habit:
    mode: str               # styles | manual | plain (no formatting) | text (pasted)
    numbering: str
    body: float
    caps_l1: bool
    bold_levels: int        # headings of depth <= this are bold
    size_bonus: list[float]  # extra size per depth 1..4
    center_l1: bool
    miss_format: float      # a heading loses its formatting
    body_bold: float        # a body paragraph is entirely bold (emphasis)
    word_lists: bool        # Word numbering vs typed "-" bullets
    bullet: str
    caps_body_short: float  # short body lines in capitals (e.g. "NB :")


def habit(rng: random.Random) -> Habit:
    mode = rng.choices(["styles", "manual", "plain", "text"], [18, 50, 10, 22])[0]
    body = rng.choice([11.0, 12.0, 12.0, 12.0, 14.0])
    return Habit(
        mode=mode,
        numbering=rng.choice(list(NUMBERINGS)),
        body=body,
        caps_l1=rng.random() < 0.7,
        bold_levels=rng.choice([1, 2, 3, 3, 4, 4]),
        size_bonus=rng.choice([[4, 2, 0, 0], [2, 1, 0, 0], [0, 0, 0, 0], [6, 4, 2, 0], [2, 2, 0, 0]]),
        center_l1=rng.random() < 0.5,
        miss_format=rng.choice([0.0, 0.03, 0.08, 0.15]),
        body_bold=rng.choice([0.0, 0.02, 0.05, 0.1]),
        word_lists=rng.random() < 0.4,
        bullet=rng.choice(["-", "•", "–", "➢", "*", "o", "✓", "►"]),
        caps_body_short=rng.random() * 0.3,
    )


class Counter4:
    def __init__(self) -> None:
        self.c = [0, 0, 0, 0, 0]

    def next(self, depth: int) -> dict:
        self.c[depth] += 1
        for k in range(depth + 1, 5):
            self.c[k] = 0
        n = self.c[depth]
        return {"n": n, "n1": self.c[depth - 1] or 1, "n2": self.c[depth - 2] if depth >= 2 else 1,
                "R": ROMAN[(n - 1) % 10], "A": "ABCDEFGH"[(n - 1) % 8], "a": "abcdefgh"[(n - 1) % 8]}


def styled_text(e: El, h: Habit, counters: Counter4, rng: random.Random, chapter_depth: int) -> str:
    """Heading text as typed by the author (numbering + case)."""
    text = e.text
    if e.kind == "section":
        rel = e.depth - chapter_depth  # 1 = first level under the chapter
        pat = NUMBERINGS[h.numbering][min(rel, 3)] if rel >= 1 else ""
        if pat:
            text = pat.format(**counters.next(min(rel, 3))) + " " + text
    if e.depth == 1 and h.caps_l1 or (e.kind == "chapitre" and rng.random() < 0.6):
        text = text.upper()
    if e.kind == "special" and rng.random() < 0.15:
        text = text.title()
    return text


def render(els: list[El], h: Habit, rng: random.Random, tmp: Path) -> list[Raw] | None:
    """Writes the messy document and reads it back with the production extractor."""
    counters = Counter4()
    chapter_depth = 1
    lines: list[tuple[El, str]] = []
    list_num = 0
    for e in els:
        if e.label in (H1, H2, H3, H4) and e.kind in ("chapitre", "partie", "special"):
            counters = Counter4()
            chapter_depth = e.depth if e.kind == "chapitre" else chapter_depth
            if e.kind == "special":
                chapter_depth = 1
        if e.label in (H1, H2, H3, H4):
            lines.append((e, styled_text(e, h, counters, rng, chapter_depth)))
        elif e.label == LIST and e.kind == "nomark":
            lines.append((e, e.text))  # no bullet, no Word list: that's the whole difficulty
        elif e.label == LIST:
            list_num = list_num + 1 if lines and lines[-1][0].label == LIST else 1
            prefix = f"{list_num}. " if e.ordered else f"{h.bullet} "
            lines.append((e, e.text if (h.word_lists and h.mode not in ("text",)) else prefix + e.text))
        elif e.label == PARA and len(e.text) < 70 and rng.random() < h.caps_body_short:
            lines.append((e, e.text.upper()))
        else:
            lines.append((e, e.text))

    if h.mode == "text":
        sep = "\n\n" if rng.random() < 0.5 else "\n"
        raws = from_text(sep.join(t for _, t in lines))
    else:
        doc = Document()
        doc.styles["Normal"].font.size = Pt(h.body)
        for e, t in lines:
            if e.label in (H1, H2, H3, H4) and h.mode == "styles" and rng.random() > h.miss_format:
                doc.add_heading(t, level=min(e.depth, 4))
                continue
            if e.label == LIST and h.word_lists and e.kind != "nomark":
                doc.add_paragraph(t, style="List Number" if e.ordered else "List Bullet")
                continue
            p = doc.add_paragraph()
            run = p.add_run(t)
            if e.label in (H1, H2, H3, H4) and h.mode in ("manual", "styles") and rng.random() > h.miss_format:
                run.bold = e.depth <= h.bold_levels
                run.font.size = Pt(h.body + h.size_bonus[min(e.depth, 4) - 1])
                if e.depth == 1 and h.center_l1:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if e.depth >= 3 and rng.random() < 0.3:
                    run.italic = True
            elif e.label == PARA and rng.random() < h.body_bold:
                run.bold = True
            elif e.label == CAPTION and rng.random() < 0.5:
                run.italic = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        path = tmp / "d.docx"
        doc.save(path)
        raws = from_docx(path, tmp)
    raws = [r for r in raws if r.text.strip()]
    if len(raws) != len(lines):
        return None  # extraction merged/split something: skip rather than mislabel
    return raws


def sample(seed: int) -> tuple[list[Raw], list[int], Habit] | None:
    rng = random.Random(seed)
    els = truth_doc(rng)
    h = habit(rng)
    with tempfile.TemporaryDirectory() as d:
        raws = render(els, h, rng, Path(d))
    if raws is None:
        return None
    return raws, [e.label for e in els], h
