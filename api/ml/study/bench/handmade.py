"""Hand-made "visually formatted" student documents (no real Word structure), with the expected plan in a sidecar .truth.json.

Usage: python handmade.py <out dir>
Each document = one topic (content + plan) rendered with one author habit, the way students really type.
"""
import json, random, sys
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor

OUT = Path(sys.argv[1]); OUT.mkdir(parents=True, exist_ok=True)

S = [
    "Cette structure occupe une place importante dans le tissu économique de la région du Centre.",
    "Les données ont été collectées auprès de quarante personnes pendant une période de trois mois.",
    "Il ressort de cette analyse que les procédures ne sont pas toujours respectées par le personnel.",
    "Nous avons participé à plusieurs activités sous la supervision de notre encadreur professionnel.",
    "Le service dispose de cinq ordinateurs, de deux imprimantes et d'une connexion internet instable.",
    "Cette situation s'explique par l'insuffisance des moyens financiers et matériels mis à disposition.",
    "Selon Mbarga (2019), la qualité du service dépend avant tout de la formation continue des agents.",
    "Les résultats obtenus montrent une nette amélioration par rapport à l'année précédente.",
    "Toutefois, quelques difficultés ont été rencontrées, notamment les coupures régulières d'électricité.",
    "L'entreprise a été créée en 2004 à Douala et compte aujourd'hui plus de cent vingt employés.",
    "Ce travail nous a permis de mettre en pratique les connaissances théoriques acquises en salle.",
    "Par ailleurs, la collaboration entre les différents services reste un point à améliorer.",
    "Le tableau ci-dessous présente la répartition des dossiers traités au cours du stage.",
    "Il convient donc de proposer des solutions adaptées aux réalités du terrain.",
]


def para(rng, n=None):
    return " ".join(rng.sample(S, n or rng.randint(3, 5)))


def topic(name, rng):
    """Returns the logical document: list of (type, ...) elements."""
    P = lambda: ("p", para(rng))
    L = lambda *it: ("list", list(it))
    if name == "stage":
        cover = dict(kind="RAPPORT DE STAGE", title="LA GESTION DES DOSSIERS DE CRÉDIT À LA MICROFINANCE EXPRESS UNION DE YAOUNDÉ",
                     school="INSTITUT UNIVERSITAIRE DE LA CÔTE", author="NGONO ATEBA Marie Claire", matricule="21G00457",
                     sup="M. FOTSO Jean, Chef d'agence", year="2025-2026", period="du 1er juillet au 30 septembre 2025")
        els = [("special", "DEDICACE"), ("p", "À mes parents, pour tous les sacrifices consentis."),
               ("special", "REMERCIEMENTS"), P(), P(),
               ("special", "INTRODUCTION"), P(), P(),
               ("h", 1, "PRESENTATION DE LA STRUCTURE D'ACCUEIL", "chapitre"),
               ("h", 2, "Historique et situation géographique"), P(), P(),
               ("h", 2, "Missions et activités"), P(), L("l'octroi de crédits aux particuliers", "la collecte de l'épargne", "le transfert d'argent"), P(),
               ("h", 2, "Organisation et fonctionnement"), P(),
               ("h", 3, "La direction générale"), P(),
               ("h", 3, "Le service crédit"), P(),
               ("h", 1, "DEROULEMENT DU STAGE", "chapitre"),
               ("h", 2, "Les tâches effectuées"), P(), L("accueil des clients", "montage des dossiers", "saisie des opérations", "classement des archives"),
               ("table", "Répartition des dossiers traités", [["Mois", "Dossiers reçus", "Dossiers accordés"], ["Juillet", "42", "30"], ["Août", "55", "41"], ["Septembre", "38", "29"]]),
               ("h", 2, "Les difficultés rencontrées"), P(), P(),
               ("h", 2, "Les suggestions"), P(), L("former le personnel", "informatiser les archives"),
               ("special", "CONCLUSION"), P(),
               ("special", "BIBLIOGRAPHIE"), ("p", "Mbarga, P. (2019). La microfinance au Cameroun. Yaoundé : Presses de l'UCAC."),
               ("p", "Tchakounté, L. (2021). Gestion du crédit. Douala : Éditions Clé.")]
    elif name == "expose":
        cover = dict(kind="EXPOSÉ D'HISTOIRE", title="LES CAUSES DE LA PREMIÈRE GUERRE MONDIALE", school="UNIVERSITÉ DE DSCHANG",
                     author="Groupe 4", members=["KAMGA Paul", "NDJOCK Estelle", "ABENA Brice", "FOUDA Linda"], sup="Dr ESSOMBA", year="2025-2026")
        els = [("special", "INTRODUCTION"), P(),
               ("h", 1, "Les causes lointaines"), ("h", 2, "Les rivalités coloniales"), P(), P(),
               ("h", 2, "La course aux armements"), P(), ("h", 2, "Le système des alliances"), P(), L("la Triple Entente", "la Triple Alliance"),
               ("h", 1, "Les causes immédiates"), ("h", 2, "L'attentat de Sarajevo"), P(), ("h", 2, "L'engrenage des déclarations de guerre"), P(), P(),
               ("h", 1, "Les conséquences"), P(), L("des millions de morts", "la chute de quatre empires", "la naissance de la SDN"),
               ("special", "CONCLUSION"), P()]
    elif name == "memoire":
        cover = dict(kind="MÉMOIRE DE FIN D'ÉTUDES", title="IMPACT DES ENGRAIS ORGANIQUES SUR LE RENDEMENT DU MAÏS DANS LA LOCALITÉ DE FOUMBOT",
                     school="UNIVERSITÉ DE DSCHANG", faculty="FACULTÉ D'AGRONOMIE ET DES SCIENCES AGRICOLES", author="TCHINDA KUETE Armel", matricule="CM-UDS-20AGR1234",
                     sup="Pr NGUEFACK Julienne, Maître de Conférences", year="2025-2026", degree="Diplôme d'Ingénieur Agronome")
        els = [("special", "DEDICACE"), ("p", "À ma famille."), ("special", "REMERCIEMENTS"), P(), P(),
               ("special", "RESUME"), P(), ("special", "ABSTRACT"), ("p", "This study assesses the effect of organic fertilizers on maize yield in Foumbot. Results show a significant increase."),
               ("special", "INTRODUCTION GENERALE"), P(), P(), P(),
               ("h", 0, "CADRE THEORIQUE ET METHODOLOGIQUE", "partie"),
               ("h", 1, "REVUE DE LA LITTERATURE", "chapitre"), ("h", 2, "Généralités sur le maïs"), ("h", 3, "Origine et description"), P(), ("h", 3, "Exigences écologiques"), P(), P(),
               ("h", 2, "Les engrais organiques"), ("h", 3, "Définition"), P(), ("h", 3, "Types d'engrais organiques"), P(), L("le compost", "le fumier de volaille", "les engrais verts"),
               ("h", 1, "MATERIEL ET METHODES", "chapitre"), ("h", 2, "Présentation de la zone d'étude"), P(), ("fig", "Localisation de la zone d'étude"),
               ("h", 2, "Dispositif expérimental"), P(), ("table", "Traitements appliqués", [["Traitement", "Dose (t/ha)"], ["T0", "0"], ["T1", "5"], ["T2", "10"]]),
               ("h", 2, "Analyse des données"), P(),
               ("h", 0, "RESULTATS ET DISCUSSION", "partie"),
               ("h", 1, "PRESENTATION DES RESULTATS", "chapitre"), ("h", 2, "Effet sur la croissance"), P(), P(), ("h", 2, "Effet sur le rendement"), P(),
               ("table", "Rendement moyen par traitement", [["Traitement", "Rendement (t/ha)"], ["T0", "1,8"], ["T1", "2,9"], ["T2", "3,6"]]),
               ("h", 1, "DISCUSSION ET RECOMMANDATIONS", "chapitre"), ("h", 2, "Discussion"), P(), P(), ("h", 2, "Recommandations"), P(), L("vulgariser le compostage", "subventionner le fumier"),
               ("special", "CONCLUSION GENERALE"), P(), P(),
               ("special", "REFERENCES BIBLIOGRAPHIQUES"), ("p", "FAO (2020). Guide de fertilisation organique. Rome."), ("p", "Nguefack, J. (2018). Le maïs au Cameroun. Dschang University Press.")]
    elif name == "tp":
        cover = dict(kind="COMPTE RENDU DE TP", title="CONFIGURATION D'UN RÉSEAU LOCAL SOUS CISCO PACKET TRACER", school="INSTITUT UNIVERSITAIRE DE TECHNOLOGIE FOTSO VICTOR",
                     author="MBIDA Kevin", matricule="22I0456", sup="M. TAKAM", year="2025-2026")
        els = [("h", 1, "Objectif du TP"), P(), ("h", 1, "Matériel utilisé"), L("3 commutateurs Cisco 2960", "1 routeur 2811", "6 postes de travail", "des câbles droits et croisés"),
               ("h", 1, "Manipulations"), ("h", 2, "Création des VLAN"), P(), ("p", "Switch(config)# vlan 10"), ("p", "Switch(config-vlan)# name COMPTA"),
               ("h", 2, "Configuration du routage inter-VLAN"), P(), P(), ("h", 2, "Tests de connectivité"), P(),
               ("table", "Résultats des tests ping", [["Source", "Destination", "Résultat"], ["PC1", "PC2", "Succès"], ["PC1", "PC5", "Succès"], ["PC3", "PC6", "Échec"]]),
               ("h", 1, "Difficultés rencontrées"), P(), ("h", 1, "Conclusion"), P()]
    elif name == "cours":
        cover = None
        els = [("title", "COURS DE DROIT CIVIL : LES PERSONNES"), ("h", 1, "LA PERSONNALITE JURIDIQUE", "chapitre"), ("h", 2, "Définition"), P(),
               ("h", 2, "Le commencement de la personnalité"), P(), P(), ("h", 2, "La fin de la personnalité"), P(), L("le décès", "l'absence", "la disparition"),
               ("h", 1, "LES ATTRIBUTS DE LA PERSONNE", "chapitre"), ("h", 2, "Le nom"), ("h", 3, "L'attribution du nom"), P(), ("h", 3, "Le changement de nom"), P(),
               ("h", 2, "Le domicile"), P(), P(), ("h", 2, "La nationalité"), P(),
               ("h", 1, "LA CAPACITE JURIDIQUE", "chapitre"), ("h", 2, "La capacité de jouissance"), P(), ("h", 2, "La capacité d'exercice"), P(), P()]
    elif name == "projet":
        cover = dict(kind="RAPPORT DE PROJET TUTORÉ", title="CONCEPTION ET RÉALISATION D'UNE APPLICATION DE GESTION DE PHARMACIE", school="ÉCOLE NATIONALE SUPÉRIEURE POLYTECHNIQUE DE YAOUNDÉ",
                     author="Groupe 2", members=["EYENGA Sandrine (21P123)", "TALLA Yannick (21P088)", "OUMAROU Ali (21P201)"], sup="Dr KOUAMOU Georges", year="2025-2026")
        els = [("special", "RESUME"), P(), ("special", "INTRODUCTION"), P(), P(),
               ("h", 1, "ETUDE DE L'EXISTANT", "chapitre"), ("h", 2, "Description du système actuel"), P(), ("h", 2, "Critique de l'existant"), P(), L("lenteur des recherches", "erreurs de stock", "absence de statistiques"),
               ("h", 2, "Solution proposée"), P(),
               ("h", 1, "ANALYSE ET CONCEPTION", "chapitre"), ("h", 2, "Méthode d'analyse"), P(), ("h", 2, "Diagramme de cas d'utilisation"), P(), ("fig", "Diagramme de cas d'utilisation"),
               ("h", 2, "Diagramme de classes"), P(), ("fig", "Diagramme de classes"),
               ("h", 1, "REALISATION", "chapitre"), ("h", 2, "Outils utilisés"), L("Java 17", "MySQL 8", "NetBeans"), ("h", 2, "Présentation des interfaces"), P(), P(),
               ("special", "CONCLUSION ET PERSPECTIVES"), P(), ("special", "WEBOGRAPHIE"), ("p", "https://docs.oracle.com/javase, consulté le 12 mai 2026"), ("p", "https://dev.mysql.com/doc, consulté le 20 mai 2026")]
    elif name == "english":
        cover = dict(kind="INTERNSHIP REPORT", title="CUSTOMER SERVICE MANAGEMENT AT MTN CAMEROON BUEA BRANCH", school="UNIVERSITY OF BUEA", author="NGWA Blessing Fru", matricule="SM20A456",
                     sup="Mr. TABI Peter", year="2025-2026")
        els = [("special", "DEDICATION"), ("p", "To my beloved family."), ("special", "ACKNOWLEDGEMENTS"), ("p", "I wish to thank my supervisor for his guidance and all the staff of the Buea branch for their warm welcome."),
               ("special", "INTRODUCTION"), ("p", "This report presents the activities carried out during a two-month internship. It describes the company and the tasks performed."),
               ("h", 1, "PRESENTATION OF THE COMPANY", "chapter"), ("h", 2, "History"), ("p", "MTN Cameroon was launched in 2000 and has grown into the leading mobile operator. The Buea branch opened in 2006 and serves the South West Region."),
               ("h", 2, "Organisational structure"), ("p", "The branch is headed by a branch manager. He is assisted by a customer care supervisor and a sales supervisor."),
               ("h", 1, "ACTIVITIES CARRIED OUT", "chapter"), ("h", 2, "Reception of customers"), ("p", "Every morning we welcomed customers and directed them to the right desk. We recorded each complaint in the register."),
               ("h", 2, "SIM card registration"), ("p", "We registered new SIM cards using the customer's national identity card. More than three hundred cards were registered."), L("identification of the customer", "data entry", "activation"),
               ("h", 2, "Difficulties encountered"), ("p", "Network failures often slowed down the work. Some customers were impatient because of long queues."),
               ("special", "CONCLUSION"), ("p", "This internship was a rewarding experience. It helped us to understand customer relationship management in a telecommunications company.")]
    else:  # lettre
        cover = None
        els = [("p", "NKOULOU Serge Patrick"), ("p", "Étudiant en Licence 3 Informatique"), ("p", "Tél : 6 77 12 34 56"), ("p", "Yaoundé, le 12 octobre 2026"),
               ("p", "À Monsieur le Directeur Général de CAMTEL"), ("p", "Objet : Demande de stage académique"),
               ("p", "Monsieur le Directeur Général,"), P(), P(), ("p", "Dans l'attente d'une suite favorable, je vous prie d'agréer, Monsieur le Directeur Général, l'expression de mon profond respect.")]
    return cover, els


ROMANS = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII"]
HABITS = {
    # name: (heading look, numbering scheme, bullets, page break, cover layout, extras)
    "gras_centre":   dict(look="bold_caps_center", num="roman_dash", bullet="- ", brk="empty", cover="lines"),
    "souligne":      dict(look="underline", num="roman_dot", bullet="• ", brk="none", cover="lines"),
    "taille":        dict(look="size_only", num="decimal", bullet="* ", brk="page", cover="table"),
    "caps_seules":   dict(look="caps_plain", num="chap_word", bullet="-", brk="empty", cover="lines", soft=True),
    "tout_normal":   dict(look="none", num="roman_paren", bullet="→ ", brk="none", cover="lines_nobreak"),
    "styles_faux":   dict(look="all_h1", num="letter", bullet="word", brk="page", cover="table"),
    "couleur":       dict(look="color_bold", num="digit_dash", bullet="o ", brk="empty", cover="tabs", spaces=True),
    "retours_ligne": dict(look="bold", num="roman_dash", bullet="- ", brk="none", cover="lines", soft=True),
}


def number(scheme, level, counters, kind):
    counters[level] += 1
    for k in range(level + 1, 5): counters[k] = 0
    n = counters[level]
    if kind in ("chapitre", "chapter"):
        word = "CHAPTER" if kind == "chapter" else "CHAPITRE"
        return {"chap_word": f"{word} {ROMANS[n-1]} : ", "roman_dash": f"{word} {n} : ", "decimal": f"{word} {n}. ", "letter": f"{word} {ROMANS[n-1]} - "}.get(scheme, f"{word} {ROMANS[n-1]} : ")
    if kind == "partie":
        return ["PREMIERE", "DEUXIEME", "TROISIEME"][n - 1] + " PARTIE : "
    d = level - 1 if counters[0] or counters["chap"] else level  # depth under a chapter
    d = max(1, level - (1 if counters["chap"] else 0))
    if scheme == "decimal":
        path = [str(counters[k]) for k in range(1, level + 1) if counters[k]]
        return ".".join(path) + ". " if len(path) > 1 else path[0] + ". "
    seq = {
        "roman_dash": [f"{ROMANS[n-1]}- ", f"{n}- ", f"{chr(96+n)}) "],
        "roman_dot": [f"{ROMANS[n-1]}. ", f"{chr(64+n)}. ", f"{n}. "],
        "chap_word": [f"{ROMANS[n-1]}. ", f"{n}. ", f"{chr(96+n)}. "],
        "roman_paren": [f"{ROMANS[n-1]}) ", f"{n}) ", f"{chr(96+n)}) "],
        "letter": [f"{chr(64+n)}/ ", f"{n}/ ", f"{chr(96+n)}/ "],
        "digit_dash": [f"{n}- ", f"{n}.{1} ", f"- "],
    }[scheme]
    return seq[min(d, 3) - 1]


def build(name, hname, seed, with_cover=True, old_toc=False):
    rng = random.Random(seed)
    cover, els = topic(name, rng)
    h = HABITS[hname]
    doc = Document()
    base_font = rng.choice(["Calibri", "Times New Roman", "Arial", "Comic Sans MS"])
    doc.styles["Normal"].font.name = base_font
    doc.styles["Normal"].font.size = Pt(rng.choice([11, 12, 12, 13]))
    truth = []

    def add(text, bold=False, center=False, size=None, underline=False, color=None, italic=False, style=None, align=None):
        p = doc.add_paragraph(style=style)
        r = p.add_run(text)
        r.bold = bold or None; r.underline = underline or None; r.italic = italic or None
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor(*color)
        if center: p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if align: p.alignment = align
        return p

    def pagebreak():
        if h["brk"] == "empty":
            for _ in range(rng.randint(6, 14)): doc.add_paragraph("")
        elif h["brk"] == "page":
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    if cover and with_cover:
        c = cover
        fr = ["RÉPUBLIQUE DU CAMEROUN", "Paix – Travail – Patrie", "*********", "MINISTÈRE DE L'ENSEIGNEMENT SUPÉRIEUR", "*********", c["school"]] + ([c["faculty"]] if c.get("faculty") else [])
        en = ["REPUBLIC OF CAMEROON", "Peace – Work – Fatherland", "*********", "MINISTRY OF HIGHER EDUCATION", "*********"]
        if h["cover"] == "table":
            t = doc.add_table(rows=1, cols=2)
            t.cell(0, 0).text = "\n".join(fr); t.cell(0, 1).text = "\n".join(en)
        elif h["cover"] == "tabs":
            for a, b in zip(fr, en + [""] * 3): add(f"{a}\t\t\t{b}", bold=True, size=10)
        else:
            for line in fr: add(line, bold=line.isupper(), center=True, size=11)
        for _ in range(rng.randint(1, 3)): doc.add_paragraph("")
        add(c["kind"], bold=True, center=True, size=20)
        doc.add_paragraph("")
        add(("THÈME : " if rng.random() < 0.6 else "") + c["title"], bold=True, center=True, size=16)
        for _ in range(rng.randint(1, 3)): doc.add_paragraph("")
        if c.get("period"): add(f"Stage effectué {c['period']}", center=True, italic=True)
        if c.get("degree"): add(f"En vue de l'obtention du {c['degree']}", center=True, italic=True)
        if c.get("members"):
            add("Présenté par :" if rng.random() < 0.5 else "Membres du groupe :", bold=True, underline=True)
            for m in c["members"]: add(("- " if rng.random() < 0.5 else "") + m)
        else:
            add("Rédigé et présenté par :", bold=True, center=True)
            add(c["author"], bold=True, center=True)
            if c.get("matricule"): add(f"Matricule : {c['matricule']}", center=True)
        doc.add_paragraph("")
        add("Sous la supervision de :" if rng.random() < 0.5 else "Encadreur :", bold=True, center=True)
        add(c["sup"], center=True)
        for _ in range(rng.randint(1, 2)): doc.add_paragraph("")
        add(f"Année académique {c['year']}", bold=True, center=True)
        if h["cover"] != "lines_nobreak": pagebreak()

    if old_toc:
        add("SOMMAIRE", bold=True, center=True, size=14)
        k = 1
        for e in els:
            if e[0] in ("special", "h"):
                text = e[1] if e[0] == "special" else e[2]
                add(f"{text} {'.' * rng.randint(20, 60)} {k}"); k += rng.randint(1, 3)
        pagebreak()

    counters = {0: 0, 1: 0, 2: 0, 3: 0, 4: 0, "chap": 0}
    tn = fn = 0
    soft_buffer = None
    for e in els:
        kind = e[0]
        if kind == "title":
            add(e[1], bold=True, center=True, size=16); continue
        if kind == "special":
            if h["brk"] != "none" and rng.random() < 0.7: pagebreak()
            look = h["look"]
            text = e[1] if rng.random() < 0.8 else e[1].capitalize()
            if look == "all_h1": doc.add_heading(text, 1)
            elif look == "underline": add(text, underline=True, center=True)
            elif look in ("none", "caps_plain"): add(text, center=rng.random() < 0.5)
            else: add(text, bold=True, center=True, size=14)
            truth.append((1, e[1]))
        elif kind == "h":
            level, text, hk = e[1], e[2], (e[3] if len(e) > 3 else "")
            if hk in ("chapitre", "chapter"): counters["chap"] = 1
            if hk and h["brk"] != "none": pagebreak()
            label = number(h["num"], level, counters, hk)
            full = label + text
            look = h["look"]
            if look == "all_h1": doc.add_heading(full, 1)
            elif look == "bold_caps_center": add(full.upper() if level <= 1 else full, bold=True, center=level <= 1, size=14 if level <= 1 else None)
            elif look == "underline": add(full, underline=True, bold=level <= 1)
            elif look == "size_only": add(full, size={0: 18, 1: 16, 2: 14, 3: 13}.get(level, 13))
            elif look == "caps_plain": add(full.upper() if level <= 2 else full)
            elif look == "color_bold": add(full, bold=True, color=(0x1F, 0x4E, 0x79) if level <= 1 else (0xC0, 0x39, 0x2B))
            elif look == "bold": add(full, bold=True)
            else: add(full)
            truth.append((level, text))
        elif kind == "p":
            text = e[1]
            if h.get("spaces"): text = "      " + text.replace(". ", ".  ")
            if h.get("soft") and rng.random() < 0.6:
                # typed with Shift+Enter at the end of each visual line
                p = doc.add_paragraph(); words = text.split(); line = []
                for w in words:
                    line.append(w)
                    if len(" ".join(line)) > 70:
                        p.add_run(" ".join(line)).add_break(); line = []
                if line: p.add_run(" ".join(line))
            else:
                p = add(text)
                if rng.random() < 0.5: p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        elif kind == "list":
            for it in e[1]:
                if h["bullet"] == "word": doc.add_paragraph(it, style="List Bullet")
                else: add(("\t" if rng.random() < 0.4 else "") + h["bullet"] + it + rng.choice(["", " ;", ";", "."]))
        elif kind == "table":
            tn += 1
            cap = f"Tableau {tn} : {e[1]}"
            if rng.random() < 0.5: add(cap, bold=True, center=True)
            if hname in ("tout_normal", "retours_ligne"):
                for row in e[2]: add("\t".join(row))
            else:
                t = doc.add_table(rows=len(e[2]), cols=len(e[2][0])); t.style = "Table Grid"
                for i, row in enumerate(e[2]):
                    for j, cell in enumerate(row): t.cell(i, j).text = cell
            if rng.random() < 0.5: add(cap, italic=True, center=True)
            if rng.random() < 0.6: add("Source : nos enquêtes, 2025", italic=True)
        elif kind == "fig":
            fn += 1
            add("[ image ]", center=True)
            add(f"Figure {fn} : {e[1]}" if rng.random() < 0.7 else f"Fig {fn}. {e[1]}", center=True, italic=True)
    fname = f"hm_{name}_{hname}{'' if with_cover else '_sansgarde'}{'_sommaire' if old_toc else ''}"
    doc.save(OUT / f"{fname}.docx")
    (OUT / f"{fname}.docx.truth.json").write_text(json.dumps(truth, ensure_ascii=False))
    return fname, doc, truth


PLAN = [
    ("stage", "gras_centre", True, True), ("stage", "tout_normal", True, False), ("stage", "retours_ligne", False, False), ("stage", "couleur", True, True),
    ("expose", "souligne", True, False), ("expose", "tout_normal", False, False), ("expose", "styles_faux", True, False), ("expose", "caps_seules", True, False),
    ("memoire", "gras_centre", True, True), ("memoire", "taille", True, False), ("memoire", "caps_seules", True, True), ("memoire", "couleur", False, False), ("memoire", "souligne", True, False),
    ("tp", "taille", True, False), ("tp", "retours_ligne", False, False), ("tp", "styles_faux", True, False),
    ("cours", "gras_centre", False, False), ("cours", "souligne", False, False), ("cours", "caps_seules", False, False),
    ("projet", "couleur", True, True), ("projet", "tout_normal", True, False), ("projet", "taille", False, False), ("projet", "styles_faux", True, True),
    ("english", "gras_centre", True, False), ("english", "souligne", False, False),
    ("lettre", "tout_normal", False, False),
]
for i, (t, hb, cov, toc) in enumerate(PLAN):
    fname, doc, truth = build(t, hb, 100 + i, cov, toc)
    if i % 4 == 1:  # the same document pasted as plain text (all formatting lost)
        text = "\n".join(p.text for p in doc.paragraphs)
        (OUT / f"{fname}.paste").write_text(text)
        (OUT / f"{fname}.paste.truth.json").write_text(json.dumps(truth, ensure_ascii=False))
print(len(list(OUT.glob("hm_*"))) // 2, "documents")
